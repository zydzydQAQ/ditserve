"""Two-file migration evidence, with exact values and round-trip verification."""

from datetime import datetime, timezone
import json
import os
from pathlib import Path
import tempfile


def _rows(value, prefix=""):
    if isinstance(value, dict):
        for key, item in sorted(value.items()):
            yield from _rows(item, prefix + "/" + str(key).replace("~", "~0").replace("/", "~1"))
    elif isinstance(value, (list, tuple)) and value:
        for index, item in enumerate(value):
            yield from _rows(item, prefix + "/" + str(index))
    else:
        # JSON string values preserve large integer IDs, null, and exact float repr.
        text = json.dumps(value, ensure_ascii=False, allow_nan=False)
        for number, start in enumerate(range(0, max(1, len(text)), 30000)):
            yield (prefix or "/", number, text[start:start + 30000])


def write_report(destination, result):
    from openpyxl import Workbook, load_workbook
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    workbook = destination / "results.xlsx"
    report = destination / "REPORT.md"
    if workbook.exists() or report.exists():
        raise ValueError("acceptance output already exists; choose a new run directory")
    rows = list(_rows(result))
    if len(rows) > 1000000:
        raise ValueError("migration evidence exceeds workbook row limit")
    book = Workbook()
    sheet = book.active
    sheet.title = "evidence"
    sheet.append(("json_pointer", "chunk", "json_value"))
    for row in rows:
        sheet.append(row)
        # Literal strings must never become spreadsheet formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=destination)
    os.close(fd)
    try:
        book.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=False)
        actual = list(check["evidence"].iter_rows(min_row=2, values_only=True))
        check.close()
        if actual != rows:
            raise ValueError("migration workbook round-trip mismatch")
        os.replace(temporary, workbook)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    status = result.get("status", "unknown")
    report.write_text("# Migration acceptance\n\n"
        + "UTC: " + datetime.now(timezone.utc).isoformat() + "\n\n"
        + "Status: **" + status + "**. Full evidence: [results.xlsx](results.xlsx).\n\n"
        + "Each workbook row stores a JSON pointer, chunk index, and exact JSON scalar. "
        + "Concatenate chunks before decoding a long scalar. Failures, missing dependencies and "
        + "skipped/unverified checks remain explicit. No generated media or tensors are retained.\n\n"
        + "GPU seconds include subprocess loading, compilation, execution and cleanup. "
        + "These are environment checks, not scientific qualification, a performance benchmark, "
        + "or evidence of SP8 inference. Unselected suites remain unverified.\n")
    return dict(workbook=str(workbook), report=str(report))
