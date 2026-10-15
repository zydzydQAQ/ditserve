"""Export a read-only, metrics-only historical snapshot with exact provenance.

Source evidence is never rewritten, locked, pruned or relocated. Active groups
are protected across both results locations. Copied source snapshots retain
their original paths and bytes; identical destination files share hard links.
"""

from __future__ import annotations

from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import shutil
import stat
import tempfile
import time


AREAS = ("results", ".cache/results-work")
SOURCE_DIRS = {"sources", "source-snapshots", "sources-final", "sources-after-validation"}
CACHE_DIRS = {"__pycache__", "compiler-cache", "torchinductor", "triton-cache", ".git"}
CONTROLS = {"gpu-budget.json", "budget.json", "queue.json", "controller.json"}
MARKERS = {"ACTIVE", "RESUME_PENDING", "retention-blocked.json"}
ACTIVE_STATUSES = {"running", "waiting", "queued", "pending", "starting", "executing"}
RAW_SUFFIXES = {".npy", ".npz", ".pt", ".pth", ".pkl", ".safetensors", ".mp4", ".webm",
                ".png", ".jpg", ".jpeg", ".svg", ".log", ".jsonl"}
SMALL_LIMIT = 4 * 1024 * 1024
ARCHIVE_LIMIT = 64 * 1024 * 1024


def _real(path):
    path = Path(os.path.abspath(path))
    if path.resolve() != path:
        raise ValueError("symlink paths are not allowed: " + str(path))
    return path


def _inventory(project):
    rows = []
    for area in AREAS:
        base = project / area
        if not base.exists():
            continue
        if base.is_symlink():
            rows.append(dict(path=area, group="*", bytes=None, kind="symlink"))
            continue
        for directory, dirs, names in os.walk(base, followlinks=False):
            parent = Path(directory)
            for name in sorted(dirs + names):
                path = parent / name
                try:
                    info = path.lstat()
                except FileNotFoundError:
                    continue
                rel = path.relative_to(base)
                group = rel.parts[0] if len(rel.parts) > 1 or not stat.S_ISREG(info.st_mode) else "@root"
                row = dict(path=str(path.relative_to(project)), group=group, bytes=info.st_size)
                if stat.S_ISLNK(info.st_mode):
                    rows.append(dict(row, kind="symlink"))
                elif stat.S_ISREG(info.st_mode):
                    rows.append(dict(row, kind="file"))
                elif not stat.S_ISDIR(info.st_mode):
                    rows.append(dict(row, kind="special"))
            dirs[:] = sorted(name for name in dirs if not (parent / name).is_symlink())
    return sorted(rows, key=lambda row: row["path"])


def _identity(info):
    return info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns


def _read_file(path, *, target=None, collect=False, limit=ARCHIVE_LIMIT):
    """Hash a regular source, refusing links and writes concurrent with the read."""
    path = _real(path)
    chunks, digest = [], hashlib.sha256()
    with os.fdopen(os.open(path, os.O_RDONLY | os.O_NOFOLLOW), "rb") as stream:
        before = os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode) or before.st_size > limit:
            raise ValueError("source is not a bounded regular file")
        output = target.open("xb") if target is not None else None
        try:
            count = 0
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                count += len(block)
                if count > limit:
                    raise ValueError("source grew beyond its transfer limit")
                digest.update(block)
                if collect:
                    chunks.append(block)
                if output is not None:
                    output.write(block)
            after = os.fstat(stream.fileno())
        finally:
            if output is not None:
                output.close()
    if _identity(before) != _identity(after) or _identity(after) != _identity(path.stat()):
        raise ValueError("source changed during read")
    return dict(bytes=count, sha256=digest.hexdigest()), b"".join(chunks)


def _protection(project, rows):
    protected, controls = {}, {}
    for row in rows:
        path = Path(row["path"])
        if row["path"] in AREAS and row["kind"] != "file":
            protected.setdefault("@root", []).append("cannot inspect linked results location: " + row["path"])
            continue
        if row["kind"] == "symlink" and str(path.parent) in AREAS:
            protected.setdefault(row["group"], []).append("cannot inspect linked activity group: " + row["path"])
            continue
        if SOURCE_DIRS.intersection(path.parts) or CACHE_DIRS.intersection(path.parts):
            continue
        if path.name not in CONTROLS | MARKERS:
            continue
        reason = None
        if path.name in MARKERS:
            reason = "active or pending continuation marker: " + row["path"]
        elif row["kind"] != "file":
            reason = "unsafe activity control: " + row["path"]
        else:
            try:
                identity, raw = _read_file(project / path, collect=True, limit=SMALL_LIMIT)
                controls[row["path"]] = (row["group"], identity)
                value = json.loads(raw)
                if not isinstance(value, dict):
                    raise ValueError("control must contain a mapping")
                if value.get("active") or value.get("active_session"):
                    reason = "active ledger: " + row["path"]
                elif path.name in {"queue.json", "controller.json"} and value.get("status") in ACTIVE_STATUSES:
                    reason = "active/pending " + str(value["status"]) + " control: " + row["path"]
            except (OSError, ValueError, TypeError) as error:
                reason = "unreadable or unstable activity control: " + row["path"] + ": " + str(error)
        if reason:
            protected.setdefault(row["group"], []).append(reason)
    return protected, controls


def _selection(row):
    path = Path(row["path"])
    if row["kind"] != "file":
        return None, "symlink or special file; never followed"
    if CACHE_DIRS.intersection(path.parts) or path.suffix in {".pyc", ".lock"}:
        return None, "rebuildable cache or transient lock; retained at source"
    if path.suffix.lower() in RAW_SUFFIXES:
        return None, "raw media/array/log; dependency closure unverified; retained at source"
    if path.suffix == ".xlsx" or path.name == "data.json.gz":
        return ("archive", "existing workbook or complete historical compressed archive") if row["bytes"] <= ARCHIVE_LIMIT else (None, "archive exceeds transfer limit; retained at source")
    if path.name in {"REPORT.md", "report.json", "frozen-report.json"} and row["bytes"] <= ARCHIVE_LIMIT:
        return "report", "historical report; byte-identical copy"
    if path.suffix.lower() in {".json", ".md", ".csv", ".tsv", ".txt", ".yaml", ".yml"} and row["bytes"] <= SMALL_LIMIT:
        return "metadata", "legacy evidence/control/retention metadata; preserved without rebinding"
    if SOURCE_DIRS.intersection(path.parts) and path.suffix in {".py", ".cpp", ".hpp", ".h", ".cc", ".sh", ".patch"} and row["bytes"] <= SMALL_LIMIT:
        return "frozen_source", "historical source identity; identical destination copies deduplicated"
    return None, "unselected or oversized evidence; retained at source; dependency review pending"


def _workbook(path, tables):
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    workbook = Workbook(write_only=True)
    for name, rows in tables.items():
        if len(rows) > 1048576:
            raise ValueError("history manifest exceeds Excel row limit")
        sheet = workbook.create_sheet(name)
        for row in rows:
            cells = []
            for value in row:
                text = "" if value is None else str(value)
                if len(text.encode("utf-16-le")) // 2 > 32767:
                    raise ValueError("history manifest value exceeds Excel cell limit")
                cell = WriteOnlyCell(sheet, value=text)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
    workbook.save(path)
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        if check.sheetnames != list(tables):
            raise ValueError("history workbook sheet verification failed")
        for name, expected in tables.items():
            actual = check[name].iter_rows()
            for row in expected:
                cells = next(actual, ())
                if len(cells) != len(row) or any(
                        (cell.value or "") != ("" if value is None else str(value)) or cell.data_type == "f"
                        for cell, value in zip(cells, row)):
                    raise ValueError("history workbook cell verification failed: " + name)
            if next(actual, None) is not None:
                raise ValueError("history workbook row count mismatch")
    finally:
        check.close()


def export_history(source_project, destination):
    """Publish a new immutable transfer directory; never alter source evidence."""
    project, destination = _real(source_project), _real(destination)
    if not project.is_dir() or destination == project or destination in project.parents:
        raise ValueError("destination must be separate from the source project")
    if any(destination == project / area or project / area in destination.parents for area in AREAS):
        raise ValueError("destination cannot be inside source results or temporary result work")
    if destination.exists():
        raise ValueError("history destination already exists; choose a new snapshot directory")
    started = time.time()
    rows = _inventory(project)
    protected, controls = _protection(project, rows)
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".history-export-", dir=destination.parent))
    copied, pending, content = [], [], {}
    try:
        for row in rows:
            if row["group"] in protected or "@root" in protected:
                pending.append(dict(row, reason="protected activity group"))
                continue
            kind, reason = _selection(row)
            if kind is None:
                pending.append(dict(row, reason=reason))
                continue
            target = staging / row["path"]
            target.parent.mkdir(parents=True, exist_ok=True)
            try:
                before, _ = _read_file(project / row["path"], target=target)
                after, _ = _read_file(project / row["path"])
                transferred, _ = _read_file(target)
                if before != after or before != transferred or before["bytes"] != row["bytes"]:
                    raise ValueError("source changed during transfer or destination differs")
                key = before["sha256"], before["bytes"]
                if key in content:
                    target.unlink()
                    os.link(content[key], target)
                else:
                    content[key] = target
                copied.append(dict(row, **before, transfer_kind=kind, reason=reason))
            except (OSError, ValueError) as error:
                if target.exists():
                    target.unlink()
                pending.append(dict(row, reason="unstable/unreadable source: " + str(error)))
        # Recheck both locations: an idle ledger/queue may become active while
        # copying. Discard only our staged copies of the affected group.
        final_rows = _inventory(project)
        final_protected, final_controls = _protection(project, final_rows)
        for group, reasons in final_protected.items():
            protected.setdefault(group, []).extend(reason for reason in reasons if reason not in protected.get(group, []))
        for path in controls.keys() | final_controls.keys():
            if controls.get(path) != final_controls.get(path):
                group = (controls.get(path) or final_controls[path])[0]
                protected.setdefault(group, []).append("activity control changed during export: " + path)
        for row in copied[:]:
            if row["group"] in protected or "@root" in protected:
                (staging / row["path"]).unlink()
                copied.remove(row)
                pending.append(dict(row, reason="group became active/unstable during export"))
        unique = {(row["sha256"], row["bytes"]) for row in copied}
        counts = dict(source_files=len(rows), copied_files=len(copied),
                      copied_logical_bytes=sum(row["bytes"] for row in copied),
                      copied_unique_bytes=sum(size for _, size in unique),
                      pending_files=len(pending), pending_bytes=sum(row["bytes"] or 0 for row in pending),
                      protected_groups=len(protected))
        kinds = Counter(row["transfer_kind"] for row in copied)
        result = dict(status="exported_with_exclusions" if pending else "exported",
                      source_project=str(project), destination=str(destination),
                      resumable=False, raw_inputs_transferred=False, source_deleted_files=0,
                      started_at=started, completed_at=time.time(), counts=counts,
                      copied_kinds=dict(kinds), protected_groups=protected,
                      pending_paths=[dict(path=row["path"], reason=row["reason"], bytes=row["bytes"]) for row in pending])
        summary = [["field", "value"], ["schema", "ditserve.migration.history.v1"],
                   ["source_project", project], ["destination", destination], ["status", result["status"]],
                   ["resumable", "false"], ["source_deleted_files", "0"],
                   ["scope", "historical metrics snapshot; raw inputs remain at source; no A100 qualification or authorization inherited"],
                   ["started_at", repr(started)], ["completed_at", repr(result["completed_at"])]]
        summary.extend([key, str(value)] for key, value in counts.items())
        tables = {"summary": summary,
                  "provenance": [["project_relative_path", "bytes", "sha256", "kind", "reason"]] +
                      [[r["path"], r["bytes"], r["sha256"], r["transfer_kind"], r["reason"]] for r in copied],
                  "not_transferred": [["project_relative_path", "bytes_not_hashed", "reason"]] +
                      [[r["path"], r["bytes"], r["reason"]] for r in pending],
                  "protected": [["group_in_both_results_locations", "reason"]] +
                      [[group, reason] for group, reasons in sorted(protected.items()) for reason in reasons]}
        _workbook(staging / "results.xlsx", tables)
        report = ["# Historical migration snapshot", "",
                  "This is a non-resumable historical metrics snapshot. Original raw inputs remain at the source; "
                  "independent reevaluation requires those inputs. Export does not qualify an experiment or transfer GPU authorization.", "",
                  "Source: `" + str(project) + "`. No source file was deleted, rewritten, or moved.", "",
                  "The root `results.xlsx` contains the complete transfer manifest with SHA256, exact byte counts, "
                  "protected groups, and every observed path not transferred. Workbook cells were read back and verified.", "",
                  "Copied historical workbooks/reports, compatibility metadata, compressed archives and frozen source "
                  "retain their original contents and project-relative paths. Original absolute paths inside metadata "
                  "are intentionally unchanged and may require archival readers; they are not runnable target configuration.", "",
                  "Identical destination files use hard links to reduce repeated frozen-source storage. No destination "
                  "file is hard-linked to the source. Preserve these links when packaging the transfer directory.", ""]
        report.extend("- " + key + ": " + str(value) for key, value in counts.items())
        report.extend(["", "Protected groups:", ""])
        report.extend("- `" + group + "`: " + "; ".join(reasons) for group, reasons in sorted(protected.items()))
        report.extend(["", "Raw media, arrays, logs, compiler caches, symlinks and oversized unselected evidence "
                       "were not copied. Dependency status for omitted raw inputs remains unverified; retain them "
                       "at source until actual evaluation, continuation or investigation dependencies close.", ""])
        (staging / "REPORT.md").write_text("\n".join(report), encoding="utf-8")
        os.rename(staging, destination)
        return result
    finally:
        if staging.exists():
            shutil.rmtree(staging)
