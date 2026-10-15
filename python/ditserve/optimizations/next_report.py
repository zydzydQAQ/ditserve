"""Bounded metadata reports for the follow-on study; never prune or compact.

All qualification remains subordinate to settled supervisor evidence. Excel
contains request measurements and event summaries, never original event rows.
Importing this module does not import torch, numpy, openpyxl, or load a model.
"""

import collections
import hashlib
import json
import math
from pathlib import Path
import tempfile
import time

from ..research.common import atomic_json, digest, file_digest
from .next_protocol import PHASES, safe_paths
from .online import paired_gate, request_identity
from .cache_followup_report import REQUEST_FIELDS, _events, _metrics
from .reporting import _small, _stats, _walk, _fmt


FILES = dict(diagnosis="diagnosis.json", sana_vae="sana-candidate.json",
             sana_graph="sana-candidate.json", sana_numerics="sana-numerics.json",
             conditional="conditional.json")
TERMINAL = frozenset(("completed", "passed", "stopped"))
MAX_EVIDENCE_BYTES = 128 * 1024 * 1024


def _number(value):
    return type(value) in (int, float) and math.isfinite(value)


def _inside(path, root):
    return path == root or root in path.parents


class _Reader:
    def __init__(self):
        self.sources = {}

    def read(self, path, errors, required=False):
        path = Path(path)
        if not path.exists() and not required:
            return {}, None
        try:
            if path.resolve() != path.absolute() or path.is_symlink():
                raise ValueError("evidence symlinks are not followed")
            before = path.stat()
            if before.st_size > MAX_EVIDENCE_BYTES:
                raise ValueError("evidence exceeds explicit 128MiB reporting bound")
            raw = path.read_bytes()
            after = path.stat()
            if (before.st_ino, before.st_size, before.st_mtime_ns) != (after.st_ino, after.st_size, after.st_mtime_ns):
                raise ValueError("evidence changed while reading")
            value = json.loads(raw)
            source = dict(path=str(path), bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest())
            self.sources[source["path"]] = source
            return value, source
        except (OSError, ValueError, TypeError) as error:
            errors.append(str(path) + ": " + str(error))
            return {}, None

    def bound(self, binding, errors, label):
        if not isinstance(binding, dict) or not binding.get("path") or not binding.get("sha256"):
            errors.append(label + ": missing source binding")
            return {}
        value, source = self.read(binding["path"], errors, required=True)
        if source and source["sha256"] != binding["sha256"]:
            errors.append(label + ": source SHA256 changed")
        return value


def _integral(intervals, elapsed):
    if not _number(elapsed) or elapsed <= 0:
        return dict(maximum=None, time_weighted_mean=None, observed_window_ms=elapsed)
    changes = collections.defaultdict(int)
    area = 0.
    for start, end in intervals:
        start, end = max(0., start), min(elapsed, end)
        if end > start:
            changes[start] += 1
            changes[end] -= 1
            area += end - start
    current = maximum = 0
    for stamp in sorted(changes):
        current += changes[stamp]
        maximum = max(maximum, current)
    return dict(maximum=maximum, time_weighted_mean=area / elapsed, observed_window_ms=elapsed,
                occupancy_integral_request_ms=area)


def _pressure(row):
    elapsed, pending, inflight = row.get("elapsed_ms"), [], []
    requests = row.get("requests", [])
    invalid = 0
    for request in requests:
        if not request.get("arrived", True):
            continue
        arrival, start, finish = (request.get(key) for key in ("arrival_ms", "start_ms", "finish_ms"))
        if not _number(elapsed) or not _number(arrival):
            invalid += 1
            continue
        pending_end = start if _number(start) else elapsed
        if pending_end < arrival:
            invalid += 1
        else:
            pending.append((arrival, pending_end))
        if _number(start):
            end = finish if _number(finish) else elapsed
            if end < start:
                invalid += 1
            else:
                inflight.append((start, end))
    prepared, finished, quanta = {}, {}, collections.Counter()
    for event in row.get("events", []):
        rid, kind = event.get("request_id"), event.get("kind")
        if rid and kind == "video_prepare" and _number(event.get("end_ms")):
            prepared[rid] = event["end_ms"]
        if rid and kind == "video_finish" and _number(event.get("end_ms")):
            finished[rid] = event["end_ms"]
        if rid and kind == "video_quantum":
            quanta[rid] += 1
    intervals = [(start, finished.get(rid, elapsed)) for rid, start in prepared.items()
                 if _number(finished.get(rid, elapsed)) and finished.get(rid, elapsed) >= start]
    peaks = [r.get("service", {}).get("peak_allocated_bytes") for r in requests if r.get("kind") == "image"]
    return dict(pending=_integral(pending, elapsed), inflight=_integral(inflight, elapsed),
        invalid_intervals=invalid, unfinished_arrived_at_end=sum(r.get("arrived", True) and r.get("status") != "completed" for r in requests),
        video_prepared_intervals=_integral(intervals, elapsed), prepared_video_count=len(prepared),
        observed_video_finish_count=len(finished), video_quantum_counts=dict(quanta),
        image_request_peak_allocated_bytes=_stats([v for v in peaks if _number(v)]),
        scope="offline host observations: pending=arrival to first start; inflight=start to finish/window end, including paused requests; not concurrent GPU execution",
        video_state_scope="prepare-end to finish-end proxy only; actual release is later and lacks a direct timestamp; not a paused-state capacity experiment",
        memory_scope="image service peaks since per-request allocator reset; not whole-process peak")


def _identity(row):
    return row.get("trace_sha256"), row.get("order_index"), row.get("arm")


def _replay_rows(data, source, attempt, reader, errors):
    rows = [(row, source, "/rows/%d" % index) for index, row in enumerate(data.get("rows", []))]
    active = data.get("active_unit") or {}
    partial = data.get("active_trace_path") or active.get("partial_path")
    if partial:
        path = Path(partial).absolute()
        if not _inside(path, attempt) or path.resolve() != path:
            errors.append("partial evidence leaves its attempt")
        else:
            value, bound = reader.read(path, errors, required=True)
            if value.get("requests") and _identity(value) not in {_identity(r) for r, _, _ in rows}:
                rows.append((value, bound, ""))
    return rows


def _aggregate(rows, diagnosis=False):
    value = _metrics(rows, expected_requests=sum(len(row.get("requests", [])) for row in rows))
    if diagnosis:
        value = dict(traces=value["traces"], recorded_requests=value["recorded_requests"], images=value["images"],
                     scope="same-image counts/latencies only; no image-only versus mixed goodput comparison")
        if value["images"]:
            value["images"].pop("goodput_rps", None)
    return value


def _online_summary(data, source, attempt, reader, errors, expected, diagnosis=False):
    bound_rows = _replay_rows(data, source, attempt, reader, errors)
    rows = [row for row, _, _ in bound_rows]
    planned = len(expected)
    identities = [_identity(row) for row in rows]
    complete = bool(planned and len(rows) == planned and len(set(identities)) == planned and set(identities) == set(expected))
    for row in rows:
        requests = row.get("requests", [])
        try:
            valid = (request_identity(requests) == row.get("trace_sha256") and
                     row.get("status") == "completed" and row.get("complete") is True and
                     all(r.get("arrived", True) and r.get("status") == "completed" for r in requests))
        except (ValueError, KeyError, TypeError):
            valid = False
        complete = complete and valid
    complete = complete and data.get("status") == "completed"
    trace_rows, request_rows = [], []
    for index, (row, binding, pointer) in enumerate(bound_rows):
        ident = {key: row.get(key) for key in ("trace_sha256", "parent_trace_sha256", "order_index", "arm", "seed", "load", "pattern")}
        trace_rows.append(dict(ident, status=row.get("status"), complete=row.get("complete"),
            elapsed_ms=row.get("elapsed_ms"), request_count=len(row.get("requests", [])),
            metrics=_aggregate([row], diagnosis), pressure=_pressure(row),
            event_summary=_events(row.get("events", []), row.get("elapsed_ms")),
            attribution=_small(row.get("attribution", {})), burst_bounds=_small(row.get("burst_bounds", [])),
            source=binding, pointer=pointer))
        for rindex, request in enumerate(row.get("requests", [])):
            measured = {key: request[key] for key in REQUEST_FIELDS if key in request}
            for key in ("encode_submit_ms", "cpu_prompt_encode_ms"):
                if key in request:
                    measured[key] = request[key]
            arrived = request.get("arrived", True)
            completed = arrived and request.get("status") == "completed"
            measured.update(on_time=bool(completed and request["finish_ms"] <= request["deadline_ms"]),
                latency_ms=request["finish_ms"] - request["arrival_ms"] if completed else None,
                wait_ms=request["start_ms"] - request["arrival_ms"] if _number(request.get("start_ms")) else None,
                measured_service_ms=request["finish_ms"] - request["start_ms"] if completed else None)
            request_rows.append(dict(ident, request=measured, source=binding, pointer=pointer + "/requests/%d" % rindex))
    groups = []
    for category, key in (("pooled", None), ("pattern", "pattern"), ("seed", "seed"), ("trace", "trace_sha256")):
        values = [None] if key is None else sorted({row.get(key) for row in rows}, key=str)
        for value in values:
            for arm in sorted({row.get("arm") for row in rows}, key=str):
                members = [r for r in rows if r.get("arm") == arm and (key is None or r.get(key) == value)]
                if members:
                    groups.append(dict(group=category, value=value, arm=arm, metrics=_aggregate(members, diagnosis)))
    comparison = None
    if not diagnosis and expected:
        comparison = paired_gate(rows, expected_trace_ids=sorted({key[0] for key in expected}))
        if complete and data.get("comparison") != comparison:
            errors.append("recorded comparison differs from request-level recomputation")
    pressure_groups = []
    for trace, arm in sorted({(row.get("trace_sha256"), row.get("arm")) for row in trace_rows}, key=str):
        members = [row for row in trace_rows if (row.get("trace_sha256"), row.get("arm")) == (trace, arm)]
        group = dict(trace_sha256=trace, arm=arm, seed=members[0].get("seed"), pattern=members[0].get("pattern"),
                     replays=len(members), unfinished_arrived_at_end=sum(r["pressure"]["unfinished_arrived_at_end"] for r in members))
        for name in ("pending", "inflight", "video_prepared_intervals"):
            values = [r["pressure"][name] for r in members]
            valid = [r for r in values if _number(r.get("occupancy_integral_request_ms")) and _number(r.get("observed_window_ms"))]
            duration = sum(r["observed_window_ms"] for r in valid)
            group[name] = dict(maximum=max((r["maximum"] for r in valid), default=None),
                time_weighted_mean=sum(r["occupancy_integral_request_ms"] for r in valid) / duration if duration else None,
                observed_window_ms=duration, valid_replays=len(valid))
        group["image_request_peak_allocated_bytes"] = max((r["pressure"]["image_request_peak_allocated_bytes"]["maximum"]
            for r in members if r["pressure"]["image_request_peak_allocated_bytes"]["n"]), default=None)
        pressure_groups.append(group)
    return dict(matrix_complete=bool(complete), expected_replays=planned, recorded_replays=len(rows),
        expected_requests=sum(expected.values()), recorded_requests=len(request_rows),
        unrecorded_expected_requests=max(0, sum(expected.values()) - len(request_rows)),
        completed_replays=sum(r.get("status") == "completed" and r.get("complete") is True for r in rows),
        traces=trace_rows, requests=request_rows, groups=groups, pressure_groups=pressure_groups, comparison=comparison,
        source=source, active_unit=_small(data.get("active_unit")),
        scope="descriptive pooled request statistics; failed/incomplete arrivals retained; each load reported separately; no p99 claim")


def _local(data, source, config):
    rows = []
    for index, item in enumerate(data.get("rows", [])):
        row = {key: _small(item[key]) for key in ("kind", "arm", "input_id", "split", "repeat", "order_index",
               "repeat_exact", "quality", "output_sha256", "preview", "status", "error") if key in item}
        row.update(service_ms=item.get("service", {}).get("invocation_wall_ms"),
                   pipeline_ms=item.get("service", {}).get("ms"), source=source, pointer="/rows/%d" % index)
        rows.append(row)
    measured = [r for r in rows if r.get("kind") == "measurement"]
    quality = [r["quality"] for r in measured if "quality" in r]
    groups = []
    for order in (0, 1):
        for arm in sorted({r.get("arm") for r in measured}, key=str):
            costs = [r["service_ms"] for r in measured if r.get("arm") == arm and r.get("order_index") == order and _number(r.get("service_ms"))]
            groups.append(dict(order_index=order, arm=arm, service_ms=_stats(costs)))
    candidate = data.get("candidate")
    inputs = {"p%d-s%d" % (i, seed) for i, _ in enumerate(config.get("direction1", {}).get("prompts", []))
              for seed in config.get("seeds", [])}
    identities = [(r.get("order_index"), r.get("arm"), r.get("input_id"), r.get("repeat")) for r in measured]
    expected = {(order, arm, name, repeat) for order in (0, 1) for arm in ("eager", candidate)
                for name in inputs for repeat in range(3)}
    matrix_complete = len(inputs) == 8 and len(identities) == 96 and len(set(identities)) == 96 and set(identities) == expected
    gains = []
    for order in (0, 1):
        means = {group["arm"]: group["service_ms"]["mean"] for group in groups if group["order_index"] == order}
        gains.append(1-means[candidate]/means["eager"] if _number(means.get(candidate)) and _number(means.get("eager")) and means["eager"] > 0 else None)
    recorded_gains = data.get("round_speed_gains")
    gains_match = bool(isinstance(recorded_gains, list) and len(recorded_gains) == 2 and
        all(_number(a) and _number(b) and math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-12)
            for a, b in zip(gains, recorded_gains)))
    gate = bool(matrix_complete and len(quality) == 32 and all(q.get("passed") is True for q in quality)
        and all(("quality" in r) == (r.get("repeat") == 0) for r in measured)
        and all(r.get("repeat_exact") is True and _number(r.get("service_ms")) for r in measured)
        and gains_match and all(_number(gain) and gain >= .05 for gain in gains))
    return dict(rows=rows, measured_count=len(measured), expected_measured_count=96,
        full_local_matrix=matrix_complete, quality_and_speed_gate=gate, recomputed_round_speed_gains=gains,
        recorded_speed_gains_match=gains_match,
        quality_checked=len(quality), quality_passed_count=sum(q.get("passed") is True for q in quality),
        quality_failed_count=sum(q.get("passed") is False for q in quality),
        repeat_mismatches=sum(r.get("repeat_exact") is False for r in measured),
        order_arm_service=groups, recorded_round_speed_gains=data.get("round_speed_gains"),
        recorded_quality_passed=data.get("quality_passed"), recorded_performance_passed=data.get("performance_passed"),
        recorded_local_qualified=data.get("local_qualified", False), capacity=_small(data.get("capacity")),
        mechanics=_small(data.get("mechanics")), scope="full invocation wall includes validation/metadata; pipeline ms is separate; local quality and speed do not alone qualify online serving")


def _recovery(data, metadata):
    record = data.get("recovery") or {}
    evidence = metadata.get("evidence") or {}
    source = metadata.get("source") or {}
    gate = evidence.get("gate") or {}
    checks = dict(source_bound=bool(source and record.get("path") == source.get("path")
            and record.get("sha256") == source.get("sha256")),
        parent_record=record.get("status") == "passed" and record.get("passed") is True
            and record.get("gate") == gate,
        candidate=evidence.get("candidate") == data.get("candidate"),
        resume_qualified=evidence.get("status") == "passed" and evidence.get("resume_qualified") is True
            and evidence.get("eligible") is True,
        complete_gate=gate.get("passed") is True and gate.get("step_count") == 32
            and gate.get("real_sana_insertions") == 3 and bool(gate.get("checks"))
            and all(value is True for value in gate["checks"].values()),
        cleanup=evidence.get("cleanup", {}).get("passed") is True)
    return dict(passed=all(checks.values()), checks=checks, source=source or None,
        gate=_small(gate), cleanup=_small(evidence.get("cleanup")),
        scope="one fixed input, 32 steps and 3 real Sana interruptions; independent recovery regression, not full cache or online qualification")


def _numerics(data, source):
    rows = []
    for index, item in enumerate(data.get("rows", [])):
        for comparison in item.get("comparisons", []) or [{}]:
            trajectory, output = comparison.get("trajectory", {}), comparison.get("output", {})
            rows.append(dict(unit_index=item.get("unit_index"), arm=item.get("arm"), status=item.get("status"),
                input_ids=item.get("input_ids"), input_id=comparison.get("input_id"),
                comparison_status=comparison.get("status"), first_exact_drift=trajectory.get("first_exact_drift"),
                first_strict_drift=trajectory.get("first_strict_drift"), all_exact=trajectory.get("all_exact"),
                all_strict=trajectory.get("all_strict"), output_quality=_small(output), error=item.get("error"),
                source=source, pointer="/rows/%d" % index))
    return dict(rows=rows, attempted_units=len(data.get("rows", [])), required_units=data.get("required_units", 12),
        completed_units=sum(r.get("status") == "completed" for r in data.get("rows", [])),
        failed_units=sum(r.get("status") == "failed" for r in data.get("rows", [])),
        diagnostic_complete=data.get("diagnostic_complete", False), performance_claims=False,
        changes_serving_default=False, scope="instrumented numerical diagnosis only; first differing boundary is not proof of the responsible kernel")


def _device_count(row, phase, config):
    """Validate the row's devices against its frozen config, never the header."""
    identity = row.get("identity") or {}
    primary, auxiliary = config.get("gpu_uuid"), config.get("v100_uuid")
    count = row.get("gpu_count")
    if not isinstance(primary, str) or not primary.startswith("GPU-") or type(count) is not int:
        return None
    if count == 1 and identity.get("gpu_uuids") == [primary] and identity.get("auxiliary_encoder", False) is False:
        return 1
    if (count == 2 and phase == "conditional" and isinstance(auxiliary, str) and auxiliary.startswith("GPU-")
            and auxiliary != primary and identity.get("gpu_uuids") == [primary, auxiliary]
            and identity.get("auxiliary_encoder") is True):
        return 2
    return None


def _gain(rows, metric, baseline, candidate):
    costs = {arm: [r.get(metric) for r in rows if r.get("arm") == arm] for arm in (baseline, candidate)}
    if any(not values or not all(_number(v) and v > 0 for v in values) for values in costs.values()):
        return None
    means = {arm: sum(values) / len(values) for arm, values in costs.items()}
    return 1 - means[candidate] / means[baseline]


def _same_numbers(left, right):
    return (isinstance(left, list) and len(left) == len(right)
            and all(_number(a) and _number(b) and math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-12)
                    for a, b in zip(left, right)))


def _component_local(data, source, component):
    rows = []
    for index, raw in enumerate(data.get("rows", [])):
        row = {key: _small(raw[key]) for key in ("kind", "arm", "input_id", "round", "repeat", "quality",
            "quality_status", "quality_error", "status", "error", "service", "encode_ms", "generation_ms",
            "completed", "completed_steps", "device", "artifact", "cache", "policy", "wan_backend",
            "generation_completed") if key in raw}
        row.update(source=source, pointer="/rows/%d" % index)
        rows.append(row)
    inputs = ["p%d-s%d" % (p, seed) for p in (0, 1) for seed in (101, 102)]
    pairs = [dict(_small(row), source=source, pointer="/pairs/%d" % index)
             for index, row in enumerate(data.get("pairs", []))]
    if component == "video_vae":
        measured = [r for r in rows if r.get("kind") == "measurement"]
        identities = [(r.get("input_id"), r.get("round"), r.get("arm"), r.get("repeat")) for r in measured]
        expected = [(name, order, arm, 0) for name in inputs for order, arms in
                    enumerate((("eager", "compiled"), ("compiled", "eager"))) for arm in arms]
        complete = identities == expected and [p.get("input_id") for p in pairs] == inputs
        qualities = [r["quality"] for r in measured if isinstance(r.get("quality"), dict)]
        gains = [_gain([dict(r, decode_ms=r.get("service", {}).get("ms")) for r in measured if r.get("round") == order],
                       "decode_ms", "eager", "compiled") for order in (0, 1)]
        match = _same_numbers(data.get("order_gains"), gains)
        gate = bool(complete and len(qualities) == 16 and all(q.get("passed") is True for q in qualities)
            and all(r.get("generation_completed") is True and r.get("quality_status") == "completed" for r in measured)
            and all(p.get("quality_passed") is True for p in pairs)
            and match and all(_number(g) and g >= .05 for g in gains))
        denominator = 16
    else:
        measured = rows
        identities = [(r.get("input_id"), r.get("arm")) for r in measured]
        expected = [(name, arm) for index, name in enumerate(inputs)
                    for arm in (("cpu", "gpu") if index % 2 == 0 else ("gpu", "cpu"))]
        complete = identities == expected and [p.get("input_id") for p in pairs] == inputs
        qualities = [p["quality"] for p in pairs if isinstance(p.get("quality"), dict)]
        gains = [_gain(measured, metric, "cpu", "gpu") for metric in ("encode_ms", "generation_ms")]
        match = _same_numbers([data.get("encode_gain"), data.get("generation_gain")], gains)
        gate = bool(complete and len(qualities) == 4 and all(q.get("passed") is True for q in qualities)
            and all(r.get("completed") is True and r.get("completed_steps") == 32
                    and r.get("cache") == "paired_008" and r.get("policy") == "block-4"
                    and r.get("wan_backend") == "eager"
                    and r.get("device") == ("cpu" if r.get("arm") == "cpu" else "cuda:1") for r in rows)
            and data.get("encoder_device") == "cuda:1" and data.get("encoder_dtype") == "float32"
            and data.get("payload_dtype") == "float32" and data.get("payload_device") == "cpu"
            and data.get("co_resident_full_generation") is True
            and match and all(_number(g) and g >= .05 for g in gains))
        denominator = 8
    gate = bool(gate and data.get("quality_epsilon") == .05 and data.get("minimum_local_gain") == .05)
    return dict(source=source, rows=rows, pairs=pairs, measured_count=len(measured), expected_measured_count=denominator,
        full_local_matrix=bool(complete), quality_and_speed_gate=gate, recorded_local_qualified=data.get("local_qualified", False),
        quality_checked=len(qualities), quality_passed_count=sum(q.get("passed") is True for q in qualities),
        quality_failed_count=sum(q.get("passed") is False for q in qualities),
        recomputed_gains=gains, recorded_gains_match=match, active_unit=_small(data.get("active_unit")),
        metadata=_small({k: v for k, v in data.items() if k not in ("rows", "pairs", "online_comparison")}),
        scope="independent component comparison; original quality/failed rows retained; no combined backend qualification")


def _conditional_components(data, attempt, config, workloads, reader, issues):
    result = {}
    expected = {(p["parent_trace_sha256"], order, arm): 24 for p in workloads.get("diagnosis", [])
                for order in (0, 1) for arm in ("baseline", "candidate")}
    for name, summary_name, local_name, online_name in (
            ("video_vae", "conditional-video-vae.json", "vae.json", "video-vae-online.json"),
            ("text_encoder", "conditional-text-encoder.json", "text_encoder.json", "encoder-online.json")):
        summary = data.get(name) or {}
        summary_path = attempt / summary_name
        wrapper_binding = summary.get("component_evidence")
        summary_file, summary_source = reader.read(summary_path, issues, required=bool(wrapper_binding))
        if wrapper_binding and (wrapper_binding.get("path") != str(summary_path) or not summary_source
                or wrapper_binding.get("sha256") != summary_source.get("sha256")):
            issues.append(name + ": component summary binding differs")
        payload = {k: v for k, v in summary.items() if k not in ("triggered", "component_evidence")}
        if summary and summary_file and payload != summary_file:
            issues.append(name + ": component summary differs from its native summary file")
        if not summary:
            summary = summary_file
        native_path = attempt / local_name
        record = summary.get("local_evidence")
        if record and (record.get("path") != str(native_path) or not record.get("sha256")):
            issues.append(name + ": local evidence binding leaves its fixed attempt path")
            local, local_source = {}, None
        else:
            local, local_source = reader.read(native_path, issues, required=bool(record))
            if record and local_source and local_source["sha256"] != record["sha256"]:
                issues.append(name + ": local evidence SHA256 differs")
        if not summary and not local_source:
            continue
        local_summary = _component_local(local, local_source, name)
        if summary.get("local_qualified") is True and (not record or not local_summary["quality_and_speed_gate"]
                or local.get("local_qualified") is not True or not wrapper_binding):
            issues.append(name + ": claimed local qualification lacks its full bound matrix")
        for field in ("local_qualified", "online_qualified", "online_evaluation_complete"):
            # Cleanup/wrapper failure may revoke an otherwise complete local
            # result. Revocation is legitimate evidence, never a source error.
            if summary.get(field) is True and local.get(field) is not True:
                issues.append(name + ": summary/native qualification differs: " + field)
        if name == "video_vae" and summary.get("local_qualified") is True:
            compiler = summary.get("compiler", {})
            if (summary.get("decoder_cleanup_passed") is not True or not compiler.get("compilation")
                    or not _number(compiler.get("graph_invocations")) or compiler["graph_invocations"] <= 0
                    or any(r.get("status") != "completed" for r in compiler["compilation"])):
                issues.append(name + ": compiled execution or decoder cleanup lacks evidence")
        online_path = summary.get("online_path") or local.get("online_path")
        online, online_source = {}, None
        if summary.get("online_path") and local.get("online_path") and summary["online_path"] != local["online_path"]:
            issues.append(name + ": local and summary online paths differ")
        if online_path and online_path != str(attempt / online_name):
            issues.append(name + ": online evidence leaves its fixed attempt path")
        else:
            online, online_source = reader.read(attempt / online_name, issues, required=bool(online_path))
        online_summary = _online_summary(online, online_source, attempt, reader, issues, expected)
        online_summary["recorded_evaluation_complete"] = online.get("evaluation_complete") is True
        if summary.get("online_qualified") is True and summary.get("online_evaluation_complete") is not True:
            issues.append(name + ": claimed online qualification lacks completion")
        result[name] = dict(summary=_small(summary), summary_source=summary_source, local=local_summary, online=online_summary,
            local_qualified=False, online_qualified=False, evaluation_complete=False,
            reason=summary.get("reason"), decision="未验证",
            scope="this component only; no automatic composition and no phase-level online qualification")
    return result


def _snapshot(output, phase, attempt, identity, reader, errors):
    path = output / "sources" / (phase + "-" + attempt.name) / "manifest.json"
    manifest, source = reader.read(path, errors, required=True)
    hashes = manifest.get("source_files", {})
    if (not hashes or digest(hashes) != manifest.get("source_sha256") or
        manifest.get("source_sha256") != identity.get("source_sha256") or
        manifest.get("config_sha256") != identity.get("config_sha256") or manifest.get("execution_root") != str(path.parent)):
        errors.append("frozen source manifest identity differs")
        return
    for name, sha in hashes.items():
        relative, file = Path(name), path.parent / name
        if (relative.is_absolute() or ".." in relative.parts or file.resolve() != file or not file.is_file()
            or file.suffix not in (".py", ".json") or file.stat().st_size > 32*1024*1024 or file_digest(file) != sha):
            errors.append("frozen source file changed: " + name)


def collect(output, ledger):
    raw_output, raw_ledger = Path(output).absolute(), Path(ledger).absolute()
    if raw_output.resolve() != raw_output or raw_ledger.resolve() != raw_ledger:
        raise ValueError("report paths must not traverse symlinks")
    output, ledger = safe_paths(raw_output, raw_ledger)
    finalization = output / "next-archive/finalization.json"
    if finalization.exists() or finalization.is_symlink():
        # Once deletion starts, never reinterpret missing raw logs as an empty
        # study. The immutable archive loader owns all receipt/source checks.
        from .next_retention import load_archive_report
        return load_archive_report(output, ledger)
    reader, errors = _Reader(), []
    study, unused = reader.read(output / "study.json", errors)
    budget, unused = reader.read(ledger, errors)
    workloads, unused = reader.read(output / "workloads.json", errors, required=bool(study))
    if study:
        auth = reader.bound(study.get("authorization"), errors, "authorization")
        if (auth.get("approved_by_user") is not True or auth.get("scope") != "serving_next_validation"
            or auth.get("budget_unlimited") is not True or auth.get("authorized_gpu_seconds", 1) is not None):
            errors.append("explicit unlimited authorization is inconsistent")
        if digest(workloads) != study.get("workloads_sha256"):
            errors.append("frozen workloads hash changed")
        for bound in study.get("historical_ledgers", []):
            reader.bound(bound, errors, "historical ledger")
        baseline = reader.bound(study.get("online_baseline_binding"), errors, "qualified online baseline")
        if not all(baseline.get(key) is True for key in ("archive_verified", "matrix_complete", "online_evaluation_complete", "online_qualified")):
            errors.append("inherited online baseline is not qualified archival evidence")
        for name, bound in study.get("cache_local_binding", {}).get("bindings", {}).items():
            reader.bound(bound, errors, "unchanged local cache " + name)
        if budget and (budget.get("authorization_binding") != study.get("authorization") or
                       budget.get("budget_unlimited") is not True or budget.get("limit_gpu_seconds", 1) is not None
                       or budget.get("config_sha256") != study.get("config_sha256")):
            errors.append("ledger differs from the frozen authorization/configuration")
        if budget:
            historical = {entry.get("path"): entry.get("sha256") for entry in study.get("historical_ledgers", [])}
            parents = {entry.get("path"): entry.get("sha256") for entry in budget.get("parents", {}).values()}
            costs = [entry.get("gpu_seconds") for entry in budget.get("sessions", [])]
            if historical != parents or len(historical) != 3:
                errors.append("ledger historical parent bindings differ")
            if (not all(_number(value) and value >= 0 for value in costs) or
                not _number(budget.get("used_gpu_seconds")) or
                not math.isclose(sum(costs), budget["used_gpu_seconds"], rel_tol=1e-12, abs_tol=1e-8)):
                errors.append("settled lifetime cost sum differs")
            for phase in PHASES:
                amount = sum(row["gpu_seconds"] for row in budget.get("sessions", [])
                             if row.get("stage") == phase and _number(row.get("gpu_seconds")))
                actual = budget.get("stage_used_gpu_seconds", {}).get(phase)
                if not _number(actual) or not math.isclose(actual, amount, rel_tol=1e-12, abs_tol=1e-8):
                    errors.append("settled stage cost sum differs: " + phase)
    phases, active, configs = [], budget.get("active"), {}
    for phase in PHASES:
        issues = []
        terminal, terminal_source = reader.read(output / phase / "terminal.json", issues)
        phase_active = active if (active or {}).get("stage") == phase else None
        candidates = sorted((output / phase).glob("attempt[0-9][0-9]"))
        selected = ((phase_active or {}).get("identity", {}).get("output") or terminal.get("output") or
                    (str(candidates[-1]) if candidates else None))
        item = dict(phase=phase, status="running" if phase_active else terminal.get("status", "not_started"),
                    decision="未验证", reason=terminal.get("reason", "阶段尚未结算" if selected else "尚未启动"),
                    authoritative=False, partial=True, online_qualified=False, errors=issues,
                    terminal_source=terminal_source, selected_attempt=selected, evidence=None, metadata={})
        if selected:
            attempt = Path(selected).absolute()
            if attempt.parent != output / phase or attempt.resolve() != attempt:
                issues.append("attempt is outside the current phase")
                phases.append(item)
                continue
            supervisor, unused = reader.read(attempt / "supervisor.json", issues, required=bool(terminal))
            config, unused = reader.read(attempt / "config.json", issues, required=True)
            configs[phase] = config
            attempt_study, unused = reader.read(attempt / "study.json", issues, required=True)
            if not study or attempt_study != study or digest(config) != study.get("config_sha256"):
                issues.append("attempt configuration/study differs")
            identity = (phase_active or {}).get("identity") or terminal
            _snapshot(output, phase, attempt, identity, reader, issues)
            evidence_path = Path(terminal.get("evidence_path", str(attempt / FILES.get(phase, phase + ".json"))))
            if not _inside(evidence_path, attempt):
                issues.append("native evidence path leaves its attempt")
                data, source = {}, None
            else:
                data, source = reader.read(evidence_path, issues, required=bool(terminal))
            item["evidence"] = source
            item["worker_status"] = data.get("status")
            item["worker_reason"] = data.get("reason")
            for name in ("loading", "capacity", "memory", "cache-factorial-warmup", "sana-profile", "sana-wan-resume"):
                value, binding = reader.read(attempt / (name + ".json"), issues)
                if binding:
                    item["metadata"][name] = dict(source=binding, evidence=_small(value))
            if phase in ("sana_vae", "sana_graph"):
                item["local"] = _local(data, source, config)
                item["local"]["recovery"] = _recovery(data, item["metadata"].get("sana-wan-resume", {}))
                if data.get("local_qualified") is True and (not item["local"]["quality_and_speed_gate"]
                    or data.get("quality_passed") is not True or data.get("performance_passed") is not True
                    or data.get("capacity", {}).get("passed") is not True
                    or not item["local"]["recovery"]["passed"]):
                    issues.append("claimed local qualification lacks complete quality/speed/capacity/recovery evidence")
                if data.get("online_evaluation_complete") is True and data.get("local_qualified") is not True:
                    issues.append("online evaluation lacks prerequisite local qualification")
                online, binding = reader.read(attempt / "sana-online.json", issues)
                expected = {(pair["parent_trace_sha256"], order, arm): 24 for pair in workloads.get("diagnosis", [])
                            for order in (0, 1) for arm in ("baseline", "candidate")}
                item["online"] = _online_summary(online, binding, attempt, reader, issues, expected)
            elif phase == "sana_numerics":
                item["numerics"] = _numerics(data, source)
            elif phase == "diagnosis" or phase.startswith("cache_load_"):
                if phase == "diagnosis":
                    expected = {(pair[arm]["trace_sha256"], order, arm): len(pair[arm]["requests"])
                                for pair in workloads.get("diagnosis", []) for order in (0, 1) for arm in ("image_only", "mixed")}
                else:
                    key = {"cache_load_05": "0.5", "cache_load_07": "0.7", "cache_load_09": "0.9"}[phase]
                    expected = {(trace["trace_sha256"], order, arm): 24 for trace in workloads.get("highload", {}).get(key, [])
                                for order in (0, 1) for arm in ("baseline", "candidate")}
                item["online"] = _online_summary(data, source, attempt, reader, issues, expected, phase == "diagnosis")
            else:
                item["conditional"] = _small(data)
                item["components"] = _conditional_components(data, attempt, config, workloads, reader, issues)
            charge = terminal.get("charge", {})
            sessions = [s for s in budget.get("sessions", []) if s.get("session_id") == charge.get("session_id")]
            final = terminal.get("final_occupancy", {})
            device_count = _device_count(charge, phase, config)
            charged_elapsed = charge.get("elapsed_seconds")
            cost_valid = bool(device_count and _number(charged_elapsed) and charged_elapsed >= 0
                and _number(charge.get("gpu_seconds"))
                and math.isclose(charge["gpu_seconds"], charged_elapsed * device_count, rel_tol=1e-12, abs_tol=1e-8))
            settled = bool(charge and len(sessions) == 1 and sessions[0] == charge and charge.get("status") == "completed"
                and charge.get("stage") == phase and cost_valid
                and charge.get("identity", {}).get("output") == str(attempt)
                and charge.get("identity", {}).get("run_root") == str(output)
                and charge.get("identity", {}).get("config_sha256") == study.get("config_sha256")
                and charge.get("identity", {}).get("source_sha256") == terminal.get("source_sha256")
                and charge.get("identity", {}).get("authorization_source") == study.get("authorization")
                and charge.get("identity", {}).get("workloads_sha256") == study.get("workloads_sha256"))
            authority = bool(not phase_active and terminal and supervisor == terminal and settled
                and terminal.get("status") in TERMINAL and terminal.get("returncode") == 0
                and not terminal.get("supervisor_failure") and source
                and final.get("status") == "exclusive_observed" and final.get("owned") == [] and final.get("foreign") == []
                and not issues and not errors)
            item.update(authoritative=authority, partial=not authority, charge=_small(charge))
            online = item.get("online")
            if online:
                comparison = online.get("comparison")
                evaluated = bool(authority and online["matrix_complete"] and
                    (data.get("online_evaluation_complete") is True if phase.startswith("sana_") else data.get("evaluation_complete") is True))
                item["evaluation_complete"] = evaluated
                item["online_qualified"] = bool(evaluated and comparison and comparison.get("passed")
                                                and terminal.get("online_qualified") is True)
                if evaluated:
                    item["decision"] = ("诊断完成" if phase == "diagnosis" else "值得继续" if item["online_qualified"] else "当前配置未过在线门槛")
            if authority and phase in ("sana_vae", "sana_graph") and not item.get("evaluation_complete"):
                if item["local"]["quality_failed_count"] or item["local"]["repeat_mismatches"]:
                    item["decision"] = "未过局部质量门槛；在线未验证"
                elif data.get("performance_passed") is False:
                    item["decision"] = "未过双序局部速度门槛；在线未验证"
            if authority and phase == "sana_numerics":
                item["decision"] = "数值诊断完成；没有性能资格" if data.get("diagnostic_complete") else "数值诊断未完成"
            if authority and phase == "conditional":
                for name, component in item.get("components", {}).items():
                    summary, local, online = component["summary"], component["local"], component["online"]
                    component["local_qualified"] = bool(summary.get("local_qualified") is True and local["quality_and_speed_gate"])
                    component["evaluation_complete"] = bool(component["local_qualified"] and online["matrix_complete"]
                        and online["recorded_evaluation_complete"] and summary.get("online_evaluation_complete") is True)
                    component["online_qualified"] = bool(component["evaluation_complete"]
                        and summary.get("online_qualified") is True and (online.get("comparison") or {}).get("passed") is True
                        and (name != "text_encoder" or device_count == 2))
                    component["decision"] = ("值得继续（该组件独立对照）" if component["online_qualified"] else
                        "当前组件未过在线门槛" if component["evaluation_complete"] else
                        "未过局部质量门槛；在线未验证" if local["quality_failed_count"] else
                        "未过局部速度门槛；在线未验证" if local["full_local_matrix"] and not local["quality_and_speed_gate"] else "未验证")
                item["decision"] = "条件组件分别记录；组合未验证" if item.get("components") else "条件分支未验证"
        phases.append(item)
    elapsed = max(0., time.monotonic() - active["started_monotonic"]) if active and _number(active.get("started_monotonic")) else 0.
    active_count = _device_count(active, active.get("stage"), configs.get(active.get("stage"), {})) if active else None
    if active and active_count is None:
        errors.append("active device/count identity differs from its frozen phase configuration")
    # A malformed row never receives an invented count. Normal active sessions
    # include exact count=1 or the explicitly admitted conditional count=2.
    active_cost = elapsed * active_count if active_count else None if active else 0.
    return dict(schema_version="ditserve.optimizations.next_report.v1", output=str(output), budget_path=str(ledger),
        status="running" if active else "not_started" if not study else "recorded",
        partial=bool(active or any(p["partial"] for p in phases) or errors), phases=phases, errors=errors,
        budget=dict(budget_unlimited=budget.get("budget_unlimited", study.get("budget_unlimited")),
            authorization_verified=bool(study and not errors), limit_gpu_seconds=budget.get("limit_gpu_seconds"),
            settled_gpu_seconds=budget.get("used_gpu_seconds", 0.), active_gpu_seconds_estimate=active_cost,
            active_device_count=active_count, active_wall_seconds_estimate=elapsed,
            active_unsettled=bool(active), total_gpu_seconds_including_active=(budget.get("used_gpu_seconds", 0.) + active_cost
                if active_cost is not None else None),
            stage_used_gpu_seconds=budget.get("stage_used_gpu_seconds", {}), sessions=_small(budget.get("sessions", [])),
            scope="row-count lifetime accounting: primary A100, conditional may explicitly hold A100+V100; full load/warmup/compile/failure/cleanup wall multiplied by held count; active is unsettled"),
        source_bindings=list(reader.sources.values()), raw_events_in_excel=False, originals_deleted=0,
        scope="finite follow-on study; local diagnostics and each cache load remain separate; no cross-load pooled qualification")


def render(report):
    output = Path(report["output"])
    lines = ["# 后续 Serving 验证", "", "状态：%s；存在未结算或未验证阶段：%s。" % (report["status"], report["partial"]), "",
        "| 阶段 | 状态 | 监督证据有效 | 结论 | 范围/原因 |", "| --- | --- | --- | --- | --- |"]
    if report.get("archive_only"):
        lines.insert(2, "归档只读：展示已核验的历史指标；原始日志已按清单收束，不能重新评测，也不授予新的资格。\n")
    for phase in report["phases"]:
        lines.append("| %s | %s | %s | %s | %s |" % tuple(str(phase[key]).replace("|", "/").replace("\n", " ") for key in ("phase", "status", "authoritative", "decision", "reason")))
        if phase.get("errors"):
            lines.append("证据问题：" + "; ".join(phase["errors"]))
    for phase in report["phases"]:
        online = phase.get("online")
        if online and online["recorded_replays"]:
            lines += ["", "**%s**：重放 %d/%d；请求 %d/%d；完整矩阵 %s。" % (phase["phase"], online["recorded_replays"],
                online["expected_replays"], online["recorded_requests"], online["expected_requests"], online["matrix_complete"]), "",
                "| 分组 | 值 | 策略 | 重放 | 图像 mean/p95(ms), n | 图像按时/到达 | goodput | 视频按时/到达 |",
                "| --- | --- | --- | --- | --- | --- | --- | --- |"]
            for group in online["groups"]:
                if group["group"] == "trace":
                    continue  # Full per-trace values remain in Excel/JSON.
                metrics = group["metrics"]
                image, video, overall = (metrics.get(key) or {} for key in ("images", "videos", "all"))
                lines.append("| %s | %s | %s | %s | %s/%s, %s | %s/%s | %s | %s/%s |" % (
                    group["group"], group["value"], group["arm"], metrics["traces"], _fmt(image.get("latency_mean_ms")),
                    _fmt(image.get("latency_p95_ms")), image.get("sample_count"), image.get("on_time"), image.get("arrived"),
                    _fmt(overall.get("goodput_rps")), _fmt(video.get("on_time")), _fmt(video.get("arrived"))))
            lines.append("失败和未完成的已到达请求留在按时率分母；goodput 含 idle 和 drain。不同 load 不混合；诊断仅比较相同图像，不比较总体 goodput。")
            if phase["phase"].startswith("cache_load_"):
                lines += ["", "| seed/pattern | 策略 | 重放 | pending 最大/时间均值 | inflight 最大/时间均值 | 结束时未完成总数 | 图像请求 reset peak 最大 MiB |",
                    "| --- | --- | --- | --- | --- | --- | --- |"]
                for group in online["pressure_groups"]:
                    peak = group["image_request_peak_allocated_bytes"]
                    lines.append("| %s/%s | %s | %s | %s/%s | %s/%s | %s | %s |" % (
                        group["seed"], group["pattern"], group["arm"], group["replays"],
                        _fmt(group["pending"]["maximum"]), _fmt(group["pending"]["time_weighted_mean"]),
                        _fmt(group["inflight"]["maximum"]), _fmt(group["inflight"]["time_weighted_mean"]),
                        group["unfinished_arrived_at_end"], _fmt(peak / (1024**2) if peak is not None else None)))
        local = phase.get("local")
        if local:
            lines += ["", "%s：局部测量 %s/96；质量通过 %s/%s，失败 %s；重复不一致 %s；正反序速度增益 %s。局部通过不等于在线资格。" %
                (phase["phase"], local["measured_count"], local["quality_passed_count"], local["quality_checked"],
                 local["quality_failed_count"], local["repeat_mismatches"], local["recomputed_round_speed_gains"]),
                "速度按完整 generate 调用 wall 计量，pipeline 同步耗时单列；32 步 / 3 次真实插入的恢复回归通过：%s。" % local["recovery"]["passed"]]
        profile = phase.get("metadata", {}).get("sana-profile", {}).get("evidence", {})
        if profile:
            summary = profile.get("summary", {})
            lines += ["", "Sana 独立诊断 profile：状态 %s；transformer host enqueue 比例 %s。host/CUDA 区间可重叠，不能求和当作 GPU 总时长。" %
                (profile.get("status"), _fmt(summary.get("transformer_host_enqueue_fraction"))),
                "", "| 阶段 | host mean/p95(ms), n | CUDA mean/p95(ms), n |", "| --- | --- | --- |"]
            for name, values in summary.get("stages", {}).items():
                host, cuda = values.get("host", {}), values.get("cuda", {})
                lines.append("| %s | %s/%s, %s | %s/%s, %s |" % (name,
                    _fmt(host.get("mean")), _fmt(host.get("p95")), host.get("n"),
                    _fmt(cuda.get("mean")), _fmt(cuda.get("p95")), cuda.get("n")))
        numeric = phase.get("numerics")
        if numeric:
            lines += ["", "数值诊断：尝试 %s/%s，完成 %s、失败 %s；只定位最早可见数值差异，不给出速度或默认策略资格。" %
                (numeric["attempted_units"], numeric["required_units"], numeric["completed_units"], numeric["failed_units"])]
        for name, component in phase.get("components", {}).items():
            local, online = component["local"], component["online"]
            lines += ["", "conditional/%s：%s；局部样本 %s/%s，质量通过 %s/%s；在线重放 %s/%s、请求 %s/%s。" %
                (name, component["decision"], local["measured_count"], local["expected_measured_count"],
                 local["quality_passed_count"], local["quality_checked"], online["recorded_replays"],
                 online["expected_replays"], online["recorded_requests"], online["expected_requests"]),
                "该组件独立对照；没有自动组合 VAE 与编码器，也没有授予条件阶段整体在线资格。"]
    budget = report["budget"]
    lines += ["", "预算：%s；已结算 %s GPU 秒，当前暂计 %s GPU 秒（未结算）。加载、预热、编译、失败及清理全部计费。" %
        ("明确无限授权" if budget["budget_unlimited"] and budget["authorization_verified"] else "授权尚未核验",
         _fmt(budget["settled_gpu_seconds"]), _fmt(budget["active_gpu_seconds_estimate"])),
        "", "队列/在途请求及 video prepare→finish 区间为离线 host 观察；不是 GPU 并行度或暂停状态容量实验。显存峰值按原记录的 reset 范围解释。",
        "原始事件未进入 Excel；没有删除、compact 或改写任何原始证据。运行中汇总仅是带来源哈希的暂存视图。", "",
        "[Excel](%s) · [结构化报告](%s) · [GPU 账本](%s)" % (output / "results.xlsx", output / "report.json", report["budget_path"])]
    if report["errors"]:
        lines += ["", "来源核验问题："] + ["- " + error for error in report["errors"]]
    return "\n".join(lines) + "\n"


def _tables(report):
    tables = {"Phases": [["phase", "status", "authoritative", "decision", "reason", "partial", "online_qualified"]],
              "Metrics": [["phase", "scope", "JSON pointer", "value"]],
              "Requests": [["phase", "trace", "order", "arm", "seed", "load", "pattern", "request_id", "kind",
                            "status", "arrived", "arrival_ms", "deadline_ms", "start_ms", "finish_ms", "latency_ms",
                            "wait_ms", "measured_service_ms", "on_time", "source", "sha256", "pointer"]],
              "Sources": [["path", "sha256", "bytes"]]}
    for phase in report["phases"]:
        tables["Phases"].append([phase.get(key) for key in tables["Phases"][0]])
        compact = {key: value for key, value in phase.items() if key not in ("online", "local", "numerics", "components")}
        for name in ("online", "local", "numerics"):
            if name in phase:
                compact[name] = {key: value for key, value in phase[name].items() if key != "requests"}
        for pointer, value in _walk(compact):
            if not isinstance(value, (dict, list)):
                tables["Metrics"].append([phase["phase"], "recorded metadata; no event bodies", pointer, value])
        for row in phase.get("online", {}).get("requests", []):
            request, source = row["request"], row.get("source") or {}
            tables["Requests"].append([phase["phase"], row.get("trace_sha256"), row.get("order_index"), row.get("arm"),
                row.get("seed"), row.get("load"), row.get("pattern")] + [request.get(key) for key in tables["Requests"][0][7:19]]
                + [source.get("path"), source.get("sha256"), row["pointer"]])
        for name, component in phase.get("components", {}).items():
            label = phase["phase"] + "/" + name
            compact_component = {key: value for key, value in component.items() if key != "online"}
            compact_component["online"] = {key: value for key, value in component["online"].items() if key != "requests"}
            for pointer, value in _walk(compact_component):
                if not isinstance(value, (dict, list)):
                    tables["Metrics"].append([label, "independent component; no composition", pointer, value])
            for row in component["online"].get("requests", []):
                request, source = row["request"], row.get("source") or {}
                tables["Requests"].append([label, row.get("trace_sha256"), row.get("order_index"), row.get("arm"),
                    row.get("seed"), row.get("load"), row.get("pattern")]
                    + [request.get(key) for key in tables["Requests"][0][7:19]]
                    + [source.get("path"), source.get("sha256"), row["pointer"]])
    for pointer, value in _walk(report["budget"]):
        if not isinstance(value, (dict, list)):
            tables["Metrics"].append(["budget", "lifetime accounting", pointer, value])
    tables["Sources"] += [[source.get(key) for key in tables["Sources"][0]] for source in report["source_bindings"]]
    return tables


def _write_excel(path, tables):
    from openpyxl import Workbook, load_workbook
    workbook = Workbook(write_only=True)
    expected = {}
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name)
        expected[name] = []
        for row in rows:
            values = [None if value is None else str(value) for value in row]
            if any(value is not None and len(value.encode("utf-16-le")) // 2 > 32767 for value in values):
                raise ValueError("Excel cell exceeds bounded literal limit")
            from openpyxl.cell import WriteOnlyCell
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
            expected[name].append(tuple(values))
    workbook.save(path)
    reopened = load_workbook(path, read_only=True, data_only=False)
    try:
        if reopened.sheetnames != list(expected):
            raise ValueError("Excel sheet readback differs")
        for name, rows in expected.items():
            actual = list(reopened[name].iter_rows(values_only=True))
            # Trailing empty cells may be omitted by the reader.
            actual = [tuple(list(row) + [None] * (len(want) - len(row))) for row, want in zip(actual, rows)] if len(actual) == len(rows) else actual
            if actual != rows:
                raise ValueError("Excel readback differs: " + name)
    finally:
        reopened.close()
    return dict(status="verified", sheets=len(tables), rows_by_sheet={name: len(rows)-1 for name, rows in tables.items()},
                formulas=False, raw_events=False, verification="every cell read back against exact emitted literal values")


def write_report(output, ledger):
    report = collect(output, ledger)
    if report.get("archive_only"):
        raise ValueError("archive-only report is immutable; use collect to inspect the verified historical report")
    output = Path(report["output"])
    output.mkdir(parents=True, exist_ok=True)
    tables = _tables(report)
    with tempfile.TemporaryDirectory(prefix=".next-report-", dir=str(output)) as directory:
        temporary = Path(directory) / "results.xlsx"
        verification = _write_excel(temporary, tables)
        temporary.replace(output / "results.xlsx")
    report["excel"] = dict(path=str(output / "results.xlsx"), sha256=file_digest(output / "results.xlsx"), **verification)
    atomic_json(output / "report.json", report)
    (output / "REPORT.md").write_text(render(report), encoding="utf-8")
    atomic_json(output / "excel-verification.json", report["excel"])
    return dict(status=report["status"], partial=report["partial"], report=str(output / "report.json"),
                markdown=str(output / "REPORT.md"), excel=report["excel"], phases=len(report["phases"]),
                requests=sum(len(p.get("online", {}).get("requests", [])) +
                    sum(len(c.get("online", {}).get("requests", [])) for c in p.get("components", {}).values())
                    for p in report["phases"]), originals_deleted=0)
