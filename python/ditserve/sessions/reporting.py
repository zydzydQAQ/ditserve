"""Auditable scalar reports for real, causally dependent LLM/image sessions.

Workbook export is deliberately independent of GPU execution. Source JSON is
flattened with explicit types, request samples remain complete, and event logs
are bounded summaries. No arrays, binaries, or large JSON blobs enter Excel.
"""

from collections import Counter, defaultdict
import hashlib
import json
import math
import os
from pathlib import Path
import tempfile

from .policy import ARMS


PHASES = ("qualify", "pilot", "state")


def _read(path, default=None):
    path = Path(path)
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else default


def _num(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _stats(values):
    values = sorted(float(x) for x in values if _num(x) and x >= 0)
    if not values:
        return dict(n=0, mean=None, p95=None, max=None)
    index = (len(values) - 1) * .95
    lo, hi = math.floor(index), math.ceil(index)
    return dict(n=len(values), mean=sum(values) / len(values),
                p95=values[lo] + (values[hi] - values[lo]) * (index - lo), max=values[-1])


def _difference(row, end, start):
    a, b = row.get(end), row.get(start)
    return a - b if _num(a) and _num(b) and a >= b else None


def _ontime(row):
    return row.get("status") == "completed" and _num(row.get("finish_ms")) and _num(row.get("deadline_ms")) \
        and row["finish_ms"] <= row["deadline_ms"]


def _session_rows(record):
    groups = defaultdict(list)
    for row in record.get("rows", []):
        groups[row.get("session_id")].append(row)
    result = []
    for sid, rows in sorted(groups.items(), key=lambda item: str(item[0])):
        rounds = {row.get("round"): row for row in rows}
        first, last = rounds.get(1), rounds.get(2)
        both = len(rows) == 2 and set(rounds) == {1, 2}
        finished = both and all(row.get("status") == "completed" for row in rows)
        start = first.get("session_start_ms", first.get("release_ms")) if first else None
        finish = last.get("finish_ms") if last else None
        result.append(dict(session_id=sid, completed=finished,
                           success=bool(finished and all(_ontime(row) for row in rows)),
                           recorded_rounds=len(rows), final_deadline_ms=last.get("deadline_ms") if last else None,
                           start_ms=start, finish_ms=finish,
                           latency_ms=finish - start if finished and _num(finish) and _num(start) and finish >= start else None))
    return result


def aggregate(records, session_count=12):
    rows = [row for record in records for row in record.get("rows", [])]
    sessions = [row for record in records for row in _session_rows(record)]
    windows = [record.get("elapsed_ms") for record in records]
    elapsed = sum(windows) if windows and all(_num(value) and value > 0 for value in windows) else None
    completed = [row for row in rows if row.get("status") == "completed"]
    ontime = sum(_ontime(row) for row in rows)
    successes = sum(row["success"] for row in sessions)
    tokens = [row.get("token_count", len(row["llm_token_ids"]) if isinstance(row.get("llm_token_ids"), list) else None)
              for row in rows]
    token_total = sum(tokens) if rows and all(_num(value) and value >= 0 for value in tokens) else None
    expected = sum(record.get("session_count", session_count) for record in records)
    return dict(replays=len(records), elapsed_ms=elapsed, planned_sessions=expected, planned_rounds=2 * expected,
                recorded_rounds=len(rows), released_rounds=sum(_num(row.get("release_ms")) for row in rows),
                missing_rounds=max(0, 2 * expected - len(rows)), completed_rounds=len(completed), ontime_rounds=ontime,
                failed_or_incomplete_rounds=sum(row.get("status") != "completed" for row in rows),
                completed_sessions=sum(row["completed"] for row in sessions), successful_sessions=successes,
                round_goodput_rps=1000 * ontime / elapsed if elapsed else None,
                session_goodput_rps=1000 * successes / elapsed if elapsed else None,
                round_throughput_rps=1000 * len(completed) / elapsed if elapsed else None,
                session_throughput_rps=1000 * sum(row["completed"] for row in sessions) / elapsed if elapsed else None,
                tokens=token_total, token_throughput_tps=1000 * token_total / elapsed if elapsed and token_total is not None else None,
                ttft_ms=_stats(_difference(row, "first_token_ms", "release_ms") for row in rows),
                ttft_from_llm_start_ms=_stats(_difference(row, "first_token_ms", "llm_start_ms") for row in rows),
                token_interval_ms=_stats(value for row in rows for value in row.get("token_intervals_ms", [])),
                image_wait_ms=_stats(_difference(row, "image_start_ms", "image_ready_ms") for row in rows),
                round_e2e_ms=_stats(_difference(row, "finish_ms", "release_ms") for row in completed),
                session_e2e_ms=_stats(row["latency_ms"] for row in sessions),
                numerical_quality_passed_replays=sum((record.get("quality") or {}).get("passed") is True for record in records),
                scope="trace wall time includes exogenous idle, endogenous think time, and final drain; no subtraction of instrumentation; latency quantiles include completed samples only")


def replay_issues(record, session_count=12):
    issues = []
    if record.get("status", "completed") not in ("passed", "completed"):
        issues.append("replay has not reported completion")
    count = record.get("session_count", session_count)
    rows = record.get("rows", [])
    if len(rows) != 2 * count or len({row.get("request_id") for row in rows}) != len(rows):
        issues.append("missing or duplicate round observations")
    groups = defaultdict(list)
    for row in rows:
        groups[row.get("session_id")].append(row)
        if row.get("status") != "completed":
            issues.append("uncompleted round " + str(row.get("request_id")))
        if isinstance(row.get("quality"), dict) and row["quality"].get("passed") is False:
            issues.append("round output/reference check failed")
        for name in ("release_ms", "deadline_ms", "finish_ms"):
            if not _num(row.get(name)) or row[name] < 0:
                issues.append("invalid " + name + " for " + str(row.get("request_id")))
        if _num(row.get("finish_ms")) and _num(row.get("release_ms")) and row["finish_ms"] < row["release_ms"]:
            issues.append("finish precedes release")
    if len(groups) != count:
        issues.append("session count differs from the frozen workload")
    for sid, turns in groups.items():
        by_round = {row.get("round"): row for row in turns}
        if len(turns) != 2 or set(by_round) != {1, 2}:
            issues.append("missing or duplicate dependent round for " + str(sid))
            continue
        first, second = by_round[1], by_round[2]
        if _num(first.get("finish_ms")) and _num(second.get("release_ms")):
            think = second.get("think_ms")
            if not _num(think) or think < 0 or not math.isclose(
                    second["release_ms"], first["finish_ms"] + think, rel_tol=0, abs_tol=1e-5):
                issues.append("second round does not depend on predecessor completion for " + str(sid))
    if not _num(record.get("elapsed_ms")) or record["elapsed_ms"] <= 0:
        issues.append("missing replay wall time")
    if (record.get("quality") or {}).get("passed") is not True:
        issues.append("numerical output/reference quality not passed")
    for name in ("trace_sha256", "device_identity", "cost_model_sha256"):
        if not record.get(name):
            issues.append("missing frozen identity " + name)
    return issues


def evaluate_pilot(records, expected_seeds=(301,), loads=(.5, .9), session_count=12):
    """Require complete within-device paired groups; never pick a winner."""
    expected = {(load, seed, arm) for load in loads for seed in expected_seeds for arm in ARMS}
    indexed, issues = {}, []
    for record in records:
        key = (record.get("load"), record.get("seed"), record.get("arm"))
        if key in indexed:
            issues.append(dict(key=key, reason="duplicate replay"))
        indexed[key] = record
        for reason in replay_issues(record, session_count):
            issues.append(dict(key=key, reason=reason))
    if len(records) != len(expected) or set(indexed) != expected:
        issues.append(dict(reason="the predeclared pilot matrix is incomplete or contains unexpected replays"))
    paired, summaries = [], []
    for load in loads:
        for seed in expected_seeds:
            group = [indexed.get((load, seed, arm)) for arm in ARMS]
            if any(record is None for record in group):
                continue
            identities = {(json.dumps(record.get("device_identity"), sort_keys=True), record.get("cost_model_sha256"),
                           record.get("trace_sha256")) for record in group}
            if len(identities) != 1:
                issues.append(dict(load=load, seed=seed, reason="paired external trace, device, or calibration identity differs"))
                continue
            if any(replay_issues(record, session_count) for record in group):
                continue
            paired.append(dict(load=load, seed=seed, arms=list(ARMS)))
            summaries.extend(dict(load=load, seed=seed, arm=record["arm"], **aggregate([record], session_count)) for record in group)
    complete = not issues and len(paired) == len(loads) * len(expected_seeds)
    return dict(status="complete" if complete else "incomplete", technical_complete=complete,
                paired_groups=paired, metrics=summaries, issues=issues, selected_candidate=None,
                interpretation="single-seed descriptive comparison; no automatic performance gate or winner; endogenous round-2 releases may differ")


def build_report(root, ledger=None):
    root = Path(root).resolve()
    if archived(root):
        return load_archive(root)
    protocol = _read(root / "protocol.json", {})
    config = _read(root / "config.json", protocol.get("config", {}))
    if ledger is None:
        ledger = protocol.get("ledger_path")
    phases = {phase: _read(root / phase / "result.json", {}) for phase in PHASES}
    budget = _read(ledger, {}) if ledger is not None else {}
    records = phases["pilot"].get("records", [])
    count = config.get("session_count", 12)
    pilot = evaluate_pilot(records, session_count=count)
    if phases["pilot"].get("status") not in ("passed", "completed"):
        pilot["technical_complete"] = False
        pilot["status"] = "incomplete"
        pilot["issues"].append(dict(reason="pilot phase has not reported technical completion"))
    for phase in PHASES:
        supervisor = _read(root / phase / "supervisor.json", {})
        if phase == "pilot" and supervisor and (supervisor.get("status") not in ("passed", "completed")
                or supervisor.get("returncode", 0) != 0 or supervisor.get("supervisor_failure")):
            pilot["technical_complete"] = False
            pilot["status"] = "incomplete"
            pilot["issues"].append(dict(reason="pilot supervision is not complete"))
    return dict(schema_version="ditserve.sessions.report.v1", root=str(root), config=config, phases=phases,
                budget=budget, budget_path=str(Path(ledger).resolve()) if ledger else None,
                pilot=pilot, descriptive_metrics=[dict(load=record.get("load"), seed=record.get("seed"), arm=record.get("arm"),
                                                      **aggregate([record], count)) for record in records],
                quality_scope="numerical/token reference checks are distinct from semantic image or conversation quality; report original checks verbatim; no unmeasured semantic quality claim",
                session_success_rule="both rounds completed by their own release-relative deadlines; final round deadline alone is insufficient",
                trace_scope="shared external starts, two user inputs, think times and seeds; second-round release follows each arm's actual first image completion")


def _hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _leaves(value, pointer=""):
    if isinstance(value, dict) and value:
        yield pointer, "dict", 0, ""
        for key in sorted(value):
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            yield from _leaves(value[key], pointer + "/" + escaped)
    elif isinstance(value, list) and value:
        yield pointer, "list", 0, ""
        for index, item in enumerate(value):
            yield from _leaves(item, pointer + "/" + str(index))
    else:
        kind = "null" if value is None else "bool" if isinstance(value, bool) else "int" if isinstance(value, int) else \
            "float" if isinstance(value, float) else "str" if isinstance(value, str) else "empty_dict" if isinstance(value, dict) else "empty_list"
        # Typed evidence uses the round-trip decimal representation of floats;
        # the human-facing round table remains numeric. Excel's numeric cell
        # precision must not alter frozen JSON when control files are retired.
        literal = value if kind in ("str", "bool") else repr(value) if kind == "float" else str(value) if kind == "int" else ""
        if kind == "float" and not math.isfinite(value):
            kind, literal = "nonfinite", repr(value)
        text = literal if isinstance(literal, str) else None
        # Excel stores <=32,767 characters. Explicit chunks avoid truncation;
        # the path/type/part tuple preserves exact reassembly and integer IDs.
        if text is not None and len(text) > 30000:
            for part in range(0, len(text), 30000):
                yield pointer, kind, part // 30000, text[part:part + 30000]
        else:
            yield pointer, kind, 0, literal


def _restore(rows):
    """Reconstruct typed scalar paths without guessing list vs numeric dict keys."""
    fragments = defaultdict(list)
    for pointer, kind, part, scalar in rows:
        fragments[(pointer or "", kind)].append((part, scalar))
    items = []
    for (pointer, kind), pieces in fragments.items():
        pieces.sort(key=lambda item: item[0])
        if [part for part, scalar in pieces] != list(range(len(pieces))):
            raise ValueError("invalid workbook scalar fragment sequence")
        raw = pieces[0][1] if len(pieces) == 1 else "".join(str(scalar or "") for part, scalar in pieces)
        if kind in ("dict", "empty_dict"):
            value = {}
        elif kind in ("list", "empty_list"):
            value = []
        elif kind == "null":
            value = None
        elif kind == "str":
            value = "" if raw is None else str(raw)
        elif kind == "int":
            value = int(raw)
        elif kind == "float":
            value = float(raw)
        elif kind == "bool" and type(raw) is bool:
            value = raw
        elif kind == "nonfinite":
            value = float(raw)
        else:
            raise ValueError("unsupported archived scalar type: " + str(kind))
        names = tuple(part.replace("~1", "/").replace("~0", "~") for part in pointer.split("/")[1:]) if pointer else ()
        items.append((names, value))
    items.sort(key=lambda item: (len(item[0]), item[0]))
    if not items or items[0][0]:
        raise ValueError("workbook is missing typed root")
    root = items[0][1]
    seen = {()}
    for names, value in items[1:]:
        if names in seen:
            raise ValueError("duplicate archived JSON pointer")
        seen.add(names)
        node = root
        for name in names[:-1]:
            node = node[int(name)] if isinstance(node, list) else node[name]
        if isinstance(node, list):
            index = int(names[-1])
            if index < 0:
                raise ValueError("negative archived list index")
            while len(node) <= index:
                node.append(None)
            node[index] = value
        elif isinstance(node, dict):
            node[names[-1]] = value
        else:
            raise ValueError("archived child below a scalar")
    return root


def load_archive(root):
    """Read the frozen workbook after control inputs have been retired.

    No evaluation is rerun, no missing raw file is treated as fresh evidence,
    and no retained result is promoted to a new scientific qualification.
    """
    from openpyxl import load_workbook
    root = Path(root).resolve()
    book = load_workbook(root / "results.xlsx", read_only=True, data_only=False)
    try:
        def restored(name):
            return _restore(book[name].iter_rows(min_row=2, values_only=True))
        value = restored("summary")
        if value.get("schema_version") != "ditserve.sessions.report.v1":
            raise ValueError("workbook does not contain a session report")
        evidence = defaultdict(list)
        for row in book["evidence"].iter_rows(min_row=2, values_only=True):
            evidence[row[0]].append(row[1:])
        value.update(config=restored("config"), budget=restored("budget"), retention=restored("retention"),
                     phases={phase: _restore(evidence[phase + "/result.json"]) if phase + "/result.json" in evidence else {}
                             for phase in PHASES}, archive_only=True)
        files = restored("provenance")
        if not isinstance(files, list):
            raise ValueError("invalid workbook provenance inventory")
        value["archive"] = dict(status="verified", archive_only=True, files=files,
                                workbook=str(root / "results.xlsx"), sha256=_hash(root / "results.xlsx"),
                                scope="previously archived metrics; no raw reevaluation")
        value["excel"] = value["archive"]
        return value
    finally:
        book.close()


def archived(root):
    root = Path(root)
    if not (root / "results.xlsx").is_file() or (root / "protocol.json").exists() or (root / "config.json").exists() \
            or any((root / phase / "result.json").exists() for phase in PHASES):
        return False
    value = load_archive(root)
    return value.get("archive_only") is True


def _update_archived_retention(root, value, retention):
    """Keep every measurement sheet byte-for-value while recording closure."""
    from openpyxl import load_workbook
    book = load_workbook(root / "results.xlsx", data_only=False)
    expected = {sheet.title: [[cell.value for cell in row] for row in sheet.iter_rows()]
                for sheet in book if sheet.title != "retention"}
    if "retention" in book:
        book.remove(book["retention"])
    _write_sheet(book, "retention", ("path", "type", "part", "value"), _leaves(retention), expected)
    handle, temp = tempfile.mkstemp(prefix=".results-", suffix=".xlsx", dir=root)
    os.close(handle)
    try:
        book.save(temp)
        book.close()
        _verify_workbook(temp, expected)
        os.replace(temp, root / "results.xlsx")
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    value["retention"] = retention
    value["archive"]["sha256"] = _hash(root / "results.xlsx")
    return value


def _event_summary(path):
    counts, status, total, bad, low, high = Counter(), Counter(), 0, 0, None, None
    def increment(counter, key):
        # Bound retained event vocabulary as well as retained payload detail.
        if key not in counter and len(counter) >= 255:
            key = "__other_categories__"
        counter[key] += 1
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            total += 1
            try:
                event = json.loads(line)
                if not isinstance(event, dict):
                    raise ValueError("event must be object")
            except (ValueError, TypeError):
                bad += 1
                continue
            increment(counts, str(event.get("kind", event.get("event", event.get("action", "unknown"))))[:500])
            increment(status, str(event.get("status", "unspecified"))[:500])
            for key in ("time_ms", "start_ms", "end_ms", "at_ms"):
                if _num(event.get(key)):
                    low = event[key] if low is None else min(low, event[key])
                    high = event[key] if high is None else max(high, event[key])
    return dict(lines=total, invalid_lines=bad, kinds=dict(counts), statuses=dict(status),
                min_time_ms=low, max_time_ms=high, aggregated_only=True,
                scope="event counts and timestamp extent; event ordering/payload is not retained and cannot be reevaluated")


def _inventory(root, value, retention):
    files, evidence, events = [], [], []
    phases = value["config"].get("phases", PHASES)
    def phase_closed(phase):
        result = value["phases"].get(phase, {})
        if (root / phase / "cleanup-unverified.json").exists() or result.get("technical_complete") is False:
            return False
        if result.get("status") == "passed":
            return True
        if phase != "state" or result.get("status") != "stopped" or result.get("technical_complete") is not True:
            return False
        supervisor = _read(root / phase / "supervisor.json", {})
        occupancy = supervisor.get("final_occupancy")
        return (supervisor.get("status") == "stopped" and supervisor.get("technical_complete") is True
                and supervisor.get("returncode") == 0 and not supervisor.get("supervisor_failure")
                and not supervisor.get("watchdog_stopped") and not supervisor.get("cleanup_pending")
                and (occupancy is None or (occupancy.get("status") == "exclusive_observed"
                                           and not occupancy.get("owned") and not occupancy.get("foreign"))))
    closed = all(phase_closed(phase) for phase in phases)
    closed = closed and value["pilot"]["technical_complete"]
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name in ("results.xlsx", "REPORT.md") or path.name.startswith(".results-"):
            continue
        relative = path.relative_to(root).as_posix()
        if path.is_symlink() or any(parent.is_symlink() for parent in path.parents if parent != root.parent):
            files.append(dict(path=relative, coverage="symlink protected", cleanup_eligible=False))
            continue
        item = dict(path=relative, bytes=path.stat().st_size, sha256=_hash(path), coverage="hash_only", cleanup_eligible=False)
        if path.suffix == ".json" and item["bytes"] <= 8 * 1024 * 1024:
            try:
                data = _read(path)
                evidence.extend((relative, *leaf) for leaf in _leaves(data))
                item["coverage"] = "complete_typed_scalars"
            except (ValueError, UnicodeError) as error:
                item["error"] = str(error)
        elif path.suffix == ".jsonl":
            try:
                summary = _event_summary(path)
                events.extend((relative, *leaf) for leaf in _leaves(summary))
                item.update(coverage="aggregated_events", cleanup_eligible=summary["invalid_lines"] == 0)
            except (ValueError, UnicodeError) as error:
                item["error"] = str(error)
        elif path.suffix == ".npy":
            item["coverage"] = "hash_only_array; scalar quality/reference evidence remains in per-round records"
            item["cleanup_eligible"] = bool(closed)
        files.append(item)
    # Preserve archived event summaries across the final refresh after raw
    # deletion. The old workbook must bind the same source bytes and digest;
    # absence alone is never interpreted as successful cleanup.
    deleted = [item for item in (retention or {}).get("files", [])
               if item.get("state") == "deleted" and not (root / item["path"]).exists()]
    event_deleted = [item for item in deleted if item.get("coverage") == "aggregated_events"]
    if event_deleted:
        from openpyxl import load_workbook
        book = load_workbook(root / "results.xlsx", read_only=True, data_only=False)
        try:
            prior = defaultdict(dict)
            for pointer, kind, part, scalar in book["provenance"].iter_rows(min_row=2, values_only=True):
                names = (pointer or "").split("/")
                if len(names) == 3 and part == 0:
                    prior[names[1]][names[2]] = int(scalar) if kind == "int" else scalar
            indexed = {item.get("path"): item for item in prior.values()}
            allowed = set()
            for item in event_deleted:
                old = indexed.get(item["path"], {})
                if old.get("sha256") != item.get("sha256") or old.get("bytes") != item.get("bytes"):
                    raise ValueError("deleted event source does not match prior workbook: " + item["path"])
                allowed.add(item["path"])
            recovered = [tuple(row) for row in book["events"].iter_rows(min_row=2, values_only=True) if row[0] in allowed]
            if {row[0] for row in recovered} != allowed:
                raise ValueError("prior workbook lacks a deleted event summary")
            events.extend(recovered)
        finally:
            book.close()
    files.extend(dict(item, cleanup_eligible=False) for item in deleted)
    return files, evidence, events


def _write_sheet(book, title, headers, rows, expected):
    sheet = book.create_sheet(title)
    matrix = [list(headers), *[list(row) for row in rows]]
    if len(matrix) > 1048576 or any(len(row) > 16384 for row in matrix):
        raise ValueError("Excel sheet limit exceeded: " + title)
    for row_index, row in enumerate(matrix, start=1):
        sheet.append(row)
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, str):
                if len(value) > 32767:
                    raise ValueError("unfragmented Excel string exceeds cell limit")
                # Worksheet bounds scan every populated cell. Address the
                # appended cells directly so export stays linear in its size.
                sheet.cell(row=row_index, column=column_index).data_type = "s"  # Includes =,+,-,@.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    expected[title] = matrix


def _verify_workbook(path, expected):
    from openpyxl import load_workbook
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        if book.sheetnames != list(expected):
            raise ValueError("workbook sheet mismatch")
        for name, matrix in expected.items():
            rows = list(book[name].iter_rows())
            if len(rows) != len(matrix):
                raise ValueError("workbook row count changed: " + name)
            for actual, planned in zip(rows, matrix):
                for cell, reference in zip(actual, planned):
                    observed = cell.value
                    if isinstance(reference, str):
                        if observed != (reference or None) or (reference and cell.data_type != "s"):
                            raise ValueError("workbook literal string changed: " + name)
                    elif isinstance(reference, bool):
                        if type(observed) is not bool or observed != reference:
                            raise ValueError("workbook boolean changed: " + name)
                    elif isinstance(reference, (float, int)):
                        if not _num(observed) or not math.isclose(reference, observed, rel_tol=1e-13, abs_tol=1e-12):
                            raise ValueError("workbook numeric precision changed: " + name)
                    elif observed is not None:
                        raise ValueError("workbook missing value changed: " + name)
    finally:
        book.close()


def _fmt(value):
    return "未测" if value is None else "%.3f" % value


def _markdown(value, archive, retention):
    lines = ["# 真实 LLM → 图像双轮会话实验", "", "同一会话的第二轮必须等第一轮图像完成并经过固定思考时间后才能释放。各策略共享外生输入，第二轮实际到达可以不同。", "",
             "首批配对状态：**%s**；不按自动性能门槛挑选胜者。单 seed 结果仅作初步描述。" % value["pilot"]["status"], "",
             "| 负载 | 策略 | 按时轮数/计划 | 成功会话/计划 | 轮 goodput /s | 会话 goodput /s | TTFT p95 ms | token 间隔 p95 ms | 图像等待 p95 ms | 轮 E2E p95 ms |", 
             "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
    for metric in value["descriptive_metrics"]:
        lines.append("| %s | %s | %s/%s | %s/%s | %s | %s | %s | %s | %s | %s |" % (
            metric["load"], metric["arm"], metric["ontime_rounds"], metric["planned_rounds"], metric["successful_sessions"], metric["planned_sessions"],
            _fmt(metric["round_goodput_rps"]), _fmt(metric["session_goodput_rps"]), _fmt(metric["ttft_ms"]["p95"]),
            _fmt(metric["token_interval_ms"]["p95"]), _fmt(metric["image_wait_ms"]["p95"]), _fmt(metric["round_e2e_ms"]["p95"])))
    lines.extend(["", "整会话成功要求两轮分别在各自截止时间前完成；最后一轮按时但第一轮超时仍失败。会话总延迟从外生首次到达到第二轮图像完成，包含思考时间。", "",
                  "goodput 分母为完整 trace 墙钟时间（包括空闲、思考和排空）；失败、缺失、依赖阻塞均保留。p95 采用线性插值且只描述已观测样本；样本数、逐轮和逐会话指标完整保存在工作簿。", "",
                  "只有完整配对、相同设备与冻结成本/外生输入、输出参考检查通过的组有比较资格；单轮数值一致性不等于语义质量或图像质量已通过独立评审。", "",
                  "资格检查、容量、状态管理及原始 quality/reference 字段保存在 evidence 表；尚未执行或失败的阶段保留原状态。", ""])
    for phase, result in value["phases"].items():
        lines.append("- %s：%s。" % (phase, result.get("status", "未开始")))
    if value["pilot"]["issues"]:
        lines.extend(["", "配对限制："] + ["- " + str(item.get("reason")) for item in value["pilot"]["issues"]])
    lines.extend(["", "归档状态：%s。" % archive["status"], "", "[轻量工作簿](results.xlsx)", "",
                  "必要控制文件、独立账本和冻结来源仍由续跑/监督器消费时保留；报告生成不自动删除它们。"])
    if retention:
        lines.extend(["", "清理凭据已纳入工作簿。已删除原始数据仅支持查阅保存的指标，无法恢复或重新评测；未完成项保留原状态。"])
        if retention.get("final_archive") is True and retention.get("status") == "completed":
            lines.extend(["", "本 run 已结束全部控制与原始数据依赖，最终交付为 results.xlsx 和 REPORT.md；独立累计账本保留在 run 外。归档查看不会触发重跑。"])
    return "\n".join(lines) + "\n"


def _publish_markdown(root, text):
    handle, temp = tempfile.mkstemp(prefix=".report-", suffix=".md", dir=root)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            stream.write(text)
        os.replace(temp, root / "REPORT.md")
    finally:
        if Path(temp).exists():
            Path(temp).unlink()


def report(root, ledger=None, retention=None, archive_only=False):
    """Publish verified results.xlsx + REPORT.md; never delete experiment data."""
    root = Path(root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    if archive_only or archived(root):
        value = load_archive(root)
        if retention is not None:
            value = _update_archived_retention(root, value, retention)
            _publish_markdown(root, _markdown(value, value["archive"], retention))
        return value
    if retention is None:
        retention = _read(root / "cleanup-summary.json")
    value = build_report(root, ledger)
    budget = value["budget"]
    if budget.get("active") or budget.get("active_session"):
        archive = dict(status="protected", reason="active GPU accounting; heavy export deferred", files=[])
    else:
        from openpyxl import Workbook
        files, evidence, events = _inventory(root, value, retention)
        book, expected = Workbook(), {}
        book.remove(book.active)
        _write_sheet(book, "summary", ("path", "type", "part", "value"), _leaves({key: item for key, item in value.items()
                         if key not in ("phases", "config", "budget")}), expected)
        _write_sheet(book, "config", ("path", "type", "part", "value"), _leaves(value["config"]), expected)
        _write_sheet(book, "budget", ("path", "type", "part", "value"), _leaves(budget), expected)
        request_fields = ("session_id", "request_id", "round", "session_start_ms", "release_ms", "deadline_ms", "status", "finish_ms",
                          "llm_start_ms", "first_token_ms", "image_ready_ms", "image_start_ms", "image_finish_ms", "token_count")
        requests, sessions = [], []
        for record in value["phases"]["pilot"].get("records", []):
            prefix = (record.get("load"), record.get("seed"), record.get("arm"))
            requests.extend((*prefix, *(row.get(field) for field in request_fields)) for row in record.get("rows", []))
            session_fields = ("session_id", "completed", "success", "recorded_rounds", "final_deadline_ms", "start_ms", "finish_ms", "latency_ms")
            sessions.extend((*prefix, *(row.get(field) for field in session_fields)) for row in _session_rows(record))
        _write_sheet(book, "rounds", ("load", "seed", "arm", *request_fields), requests, expected)
        _write_sheet(book, "sessions", ("load", "seed", "arm", "session_id", "completed", "success", "recorded_rounds", "final_deadline_ms", "start_ms", "finish_ms", "latency_ms"), sessions, expected)
        _write_sheet(book, "evidence", ("source", "path", "type", "part", "value"), evidence, expected)
        _write_sheet(book, "events", ("source", "path", "type", "part", "value"), events, expected)
        _write_sheet(book, "provenance", ("path", "type", "part", "value"), _leaves(files), expected)
        _write_sheet(book, "retention", ("path", "type", "part", "value"), _leaves(retention), expected)
        handle, temp = tempfile.mkstemp(prefix=".results-", suffix=".xlsx", dir=root)
        os.close(handle)
        try:
            book.save(temp)
            book.close()
            _verify_workbook(temp, expected)
            os.replace(temp, root / "results.xlsx")
        finally:
            if Path(temp).exists():
                Path(temp).unlink()
        archive = dict(status="verified", files=files, workbook=str(root / "results.xlsx"),
                       sha256=_hash(root / "results.xlsx"), sheets={name: len(rows) - 1 for name, rows in expected.items()})
    text = _markdown(value, archive, retention)
    _publish_markdown(root, text)
    return dict(value, archive=archive, excel=archive, retention=retention)
