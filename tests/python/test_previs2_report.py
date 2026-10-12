"""CPU-only run02 report checks: denominators, identity, budget and real exports."""

import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from ditserve.experiments import previs2_report as report


SOURCE = "a" * 64
PROTOCOL = "b" * 64


def quality(loss=.01):
    return {"status": "ok", "finite": True, "lpips_mean": loss}


def request(identity="r0", finish=100, loss=.01):
    return {"request_id": identity, "prompt_id": "new-holdout", "prompt": "fixture",
            "seed": 0, "height": 16, "width": 16, "frames": 1, "steps": 2,
            "arrival_ms": 0, "deadline_ms": 150, "finish_ms": finish,
            "status": "completed", "quality": quality(loss)}


def online(policy, horizon=200, requests=None):
    return {"type": "online_summary", "status": "completed", "case_id": "t-" + policy,
            "device": "a100", "source_identity_sha256": SOURCE, "source_verified": True,
            "pair_id": "trace-0", "policy": policy, "load": 1.2, "pattern": "bursty", "trace_seed": 0,
            "epsilon": .05, "scope": "natural", "offered": 1, "horizon_ms": horizon,
            "requests": [request()] if requests is None else requests}


def manifest(root, status="partial", charge=2):
    return {"schema_version": "ditserve.previs2.suite.v1", "status": status, "limit_seconds": 43200,
            "used_seconds": charge, "protocol_sha256": PROTOCOL,
            "current_source_identity": {"sha256": SOURCE},
            "protocol": {"expected_cases": [{"case_id": "never-run", "type": "online_summary", "device": "a100"}]},
            "records": [{"stage": "online", "device": "a100", "status": "completed", "elapsed_seconds": charge,
                         "charge_closed": True,
                         "source_identity": {"sha256": SOURCE}, "output": str(root / "00-fixture")}]}


def write_suite(root, rows, data=None):
    data = manifest(root) if data is None else data
    (root / "manifest.json").write_text(json.dumps(data))
    stage = root / "00-fixture"
    stage.mkdir(exist_ok=True)
    environment = {"type": "environment", "status": "completed", "device": "a100", "source_identity_sha256": SOURCE}
    (stage / "records.jsonl").write_text("\n".join(json.dumps(r) for r in [environment, *rows]) + "\n")
    (stage / "outcome.json").write_text(json.dumps({"status": "completed", "stage": "online"}))
    return data


class BudgetTests(unittest.TestCase):
    def test_charges_reconcile_without_resetting_or_hiding_partial(self):
        m = manifest(Path("/fixture"))
        self.assertEqual(report.budget_audit(m)["status"], "passed")
        m["used_seconds"] += 1
        self.assertIn("used_does_not_equal_attempt_charges", report.budget_audit(m)["issues"])

    def test_running_missing_bool_and_over_limit_cannot_pass(self):
        for field, value in [("elapsed_seconds", True), ("elapsed_seconds", float("nan")), ("status", "running")]:
            m = manifest(Path("/fixture"))
            m["records"][0][field] = value
            self.assertEqual(report.budget_audit(m)["status"], "failed")
        m = manifest(Path("/fixture"), charge=43201)
        self.assertIn("budget_exceeded", report.budget_audit(m)["issues"])


class PairingTests(unittest.TestCase):
    def pair(self, baseline=None, candidate=None):
        rows = [online(report.POLICIES[0]) if baseline is None else baseline,
                online(report.POLICIES[1], horizon=100) if candidate is None else candidate]
        return report.paired_goodput([report.online_metrics(r) for r in rows])

    def test_recompute_horizon_and_ignore_forged_summary(self):
        baseline = online(report.POLICIES[0])
        baseline["goodput_rps"] = 999999
        result = self.pair(baseline)
        self.assertEqual(result["pairs"][0]["baseline_rps"], 5)
        self.assertEqual(result["groups"][0]["relative_change"], 1)

    def test_unknown_quality_and_missing_requests_are_not_zero_loss(self):
        row = online(report.POLICIES[1])
        row["requests"][0]["quality"] = {"status": "unknown"}
        row["offered"] = 2
        value = report.online_metrics(row)
        self.assertEqual(value["offered"], 2)
        self.assertEqual(value["quality_unknown"], 1)
        self.assertEqual(value["missing_request_records"], 1)
        self.assertIsNone(value["goodput_rps"])
        self.assertEqual(value["goodput_lower_bound_rps"], 0)
        self.assertEqual(self.pair(candidate=row)["pairs"][0]["status"], "unavailable")

    def test_duplicate_inputs_or_arms_and_mismatch_block_pairing(self):
        for mutation in ("input", "duplicate_input", "source", "scope"):
            row = online(report.POLICIES[1])
            if mutation == "input": row["requests"][0]["seed"] = 42
            if mutation == "duplicate_input": row["requests"] *= 2; row["offered"] = 2
            if mutation == "source": row["source_identity_sha256"] = "other"
            if mutation == "scope": row["scope"] = "controlled"
            result = self.pair(candidate=row)
            self.assertFalse(result["groups"], mutation)
        row = report.online_metrics(online(report.POLICIES[0]))
        other = report.online_metrics(online(report.POLICIES[1]))
        self.assertFalse(report.paired_goodput([row, copy.deepcopy(row), other])["groups"])

    def test_zero_baseline_ratio_unavailable_and_failed_requests_retained(self):
        row = online(report.POLICIES[0], requests=[request(finish=180)])
        result = self.pair(baseline=row)
        self.assertIsNone(result["groups"][0]["relative_change"])
        self.assertEqual(result["pairs"][0]["relative_status"], "unavailable_zero_baseline")
        row["requests"][0]["status"] = "failed"
        metric = report.online_metrics(row)
        self.assertEqual(metric["failed"], 1)
        self.assertEqual(metric["offered"], 1)
        self.assertEqual(metric["goodput_rps"], 0)

    def test_strict_30_seconds_not_relative_deadline(self):
        row = online(report.POLICIES[0], horizon=40000, requests=[request(finish=30000)])
        row["requests"][0]["deadline_ms"] = 35000
        metric = report.online_metrics(row)
        self.assertEqual(metric["deadline_on_time"], 1)
        self.assertEqual(metric["under_30s"], 0)

    def test_partial_future_requests_do_not_erase_known_offered_denominator(self):
        row = online(report.POLICIES[0], requests=[request(), dict(request("future"), status="not_arrived", arrival_ms=500)])
        row["status"] = "partial"
        value = report.online_metrics(row)
        self.assertEqual(value["offered"], 1)
        self.assertEqual(value["not_arrived"], 1)
        self.assertEqual(value["missing_request_records"], 0)
        self.assertIsNone(value["goodput_rps"])
        self.assertEqual(value["goodput_lower_bound_rps"], 5)


class EvidenceTests(unittest.TestCase):
    def test_environment_identity_cannot_be_invented_from_manifest(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            write_suite(root, [online(report.POLICIES[0]), online(report.POLICIES[1])])
            path = root / "00-fixture/records.jsonl"
            rows = [json.loads(line) for line in path.read_text().splitlines()]
            rows[0].pop("source_identity_sha256")
            path.write_text("\n".join(json.dumps(r) for r in rows) + "\n")
            result = report.summarize(report.collect(root))
            self.assertFalse(result["paired_goodput"]["groups"])
            self.assertIn("unverified_source", result["paired_goodput"]["pairs"][0]["reasons"])

    def test_truncated_row_and_plan_missing_cases_retained(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            write_suite(root, [online(report.POLICIES[0])])
            with (root / "00-fixture/records.jsonl").open("a") as stream: stream.write('{"truncated":')
            result = report.summarize(report.collect(root))
            self.assertEqual(result["coverage"]["planned_case_matrix"][0]["observed_status"], "not_run")
            self.assertEqual(result["issues"][0]["reason"], "invalid_record")
            self.assertEqual(result["attempts"][0]["outcome"]["status"], "completed")

    def test_false_natural_gate_without_online_is_valid_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            m = write_suite(root, [{"type": "gate", "status": "completed", "gate_id": "controlled",
                                   "passed": True, "scope": "controlled", "logicalcap": True}])
            gate = {"schema_version": "previs2-gate.v1", "status": "failed", "source_identity_sha256": SOURCE,
                    "protocol_sha256": PROTOCOL, "mechanism_gate": True, "natural_resource_gate": False,
                    "reasons": ["natural_capacity_not_binding"]}
            (root / "gate.json").write_text(json.dumps(gate))
            result = report.build_report(root, make_figures=False, make_workbook=False)
            self.assertEqual(result["export_status"], "completed")
            self.assertFalse(result["gates"]["natural_resource_gate"])
            self.assertEqual(result["paired_goodput"]["groups"], [])
            self.assertEqual(result["gates"]["recorded_checks"][0]["scope"], "controlled")
            gate["source_identity_sha256"] = "wrong"
            self.assertIsNone(report.gates([], gate, m)["natural_resource_gate"])

    def test_event_aliases_and_missing_cost_stay_unknown(self):
        rows = [{"type": "mechanism_case", "status": "completed", "mode": "hot", "pause_step": 2,
                 "events": [{"kind": "snapshot", "ms": 4}, {"kind": "restore", "ms": 3}],
                 "memory_before": {"allocated_bytes": 100, "reserved_bytes": 200, "free_bytes": 500}},
                {"type": "mechanism_case", "status": "completed", "mode": "cold"}]
        points, capacity, _ = report.measurement_points(rows)
        self.assertEqual(points[0]["snapshot_plus_restore_ms"], 7)
        self.assertEqual(points[0]["pause_completed_steps"], 2)
        self.assertIsNone(points[1]["snapshot_plus_restore_ms"])
        self.assertEqual(capacity[0]["measurements"]["reserved_bytes"], 200)

    def test_no_run01_write_and_no_live_array_read(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            m = write_suite(root, [])
            m["schema_version"] = "ditserve.previs.suite.v1"
            (root / "manifest.json").write_text(json.dumps(m))
            with self.assertRaisesRegex(ValueError, "run01"):
                report.build_report(root, make_figures=False, make_workbook=False)
            self.assertFalse((root / "REPORT.md").exists())
            m["schema_version"] = "ditserve.previs2.suite.v1"
            m["records"][0]["status"] = "running"
            (root / "manifest.json").write_text(json.dumps(m))
            with mock.patch("ditserve.experiments.previs_verify.verify_suite") as verify:
                with self.assertRaisesRegex(ValueError, "finish"):
                    report.build_report(root, verify_raw=True, make_figures=False, make_workbook=False)
                verify.assert_not_called()

    # Figure export needs matplotlib, which lives in the GPU venv, not the CPU tools env.
    @unittest.skipUnless(importlib.util.find_spec("matplotlib"), "matplotlib only in GPU venv")
    def test_real_fixture_array_verification_and_scientific_exports(self):
        import numpy as np
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            array = np.arange(12, dtype=np.float32).reshape(1, 1, 2, 2, 3)
            np.save(root / "fixture.npy", array)
            descriptor = {"path": str(root / "fixture.npy"), "shape": list(array.shape), "dtype": "float32",
                          "finite": True, "sha256": hashlib.sha256(array.tobytes(order="C")).hexdigest()}
            rows = [{"type": "correctness", "status": "completed", "case_id": "official-none",
                     "output": descriptor, "candidate": descriptor,
                     "latent_reference": descriptor, "latent_candidate": descriptor},
                    {"type": "quality", "status": "completed", "case_id": "test-long-tram-201-selected", "measurement_scope": "fixed",
                     "split": "test", "prompt_id": "tram", "seed": 201, "reference": descriptor,
                     "quality_scope": "whole_request_fixed_policy", "scope": "controlled",
                     "elapsed_ms": 10, "quality": quality(), "output": descriptor,
                     "state_mode": "hot", "pause_completed_steps": 2, "snapshot_ms": 1, "restore_ms": 2},
                    online(report.POLICIES[0]), online(report.POLICIES[1], horizon=100)]
            m = manifest(root)
            m["protocol"].update(prompts=[{"id": "tram", "split": "test"}, {"id": "aquarium", "split": "test"}], test_seeds=[201, 202])
            write_suite(root, rows, m)
            value = report.build_report(root, verify_raw=True, make_previews=True)
            self.assertEqual(value["raw_verification"]["status"], "verified")
            self.assertEqual(value["raw_verification"]["summary"]["physical_arrays"], 1)
            self.assertEqual(value["raw_verification"]["summary"]["correctness_pair_statuses"], {"passed": 2})
            self.assertEqual(value["export_status"], "completed")
            self.assertEqual(len(value["exports"]["figures"]["artifacts"]), 3)
            self.assertEqual(value["exports"]["previews"]["status"], "partial")
            self.assertEqual(len(value["exports"]["previews"]["artifacts"]), 1)
            self.assertEqual(value["exports"]["previews"]["groups"][1]["status"], "missing")
            self.assertEqual(value["exports"]["previews"]["groups"][1]["prompt_id"], "aquarium")
            with self.assertRaisesRegex(ValueError, "verify-raw"):
                report.build_report(root, make_previews=True, make_figures=False, make_workbook=False)
            for artifact in value["exports"]["figures"]["artifacts"]:
                self.assertTrue(Path(artifact["path"]).read_bytes().startswith(b"\x89PNG"))
            workbook = load_workbook(root / "results.xlsx", read_only=True)
            self.assertIn("trigger_funnel", workbook.sheetnames)
            self.assertEqual(workbook["pairs"].max_row, 2)
            workbook.close()
            self.assertIn("自然资源gate", (root / "REPORT.md").read_text())


if __name__ == "__main__":
    unittest.main()
