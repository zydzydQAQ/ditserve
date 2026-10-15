import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import PropertyMock, patch

from ditserve.sessions.policy import ARMS
from ditserve.sessions.reporting import _verify_workbook, _write_sheet, aggregate, archived, build_report, evaluate_pilot, load_archive, report


def record(load=.5, arm="session_serial", offset=0):
    rows = []
    for index in range(2):
        start = index * 100
        first_finish = start + 10 + offset
        for turn, release in ((1, start), (2, first_finish + 5)):
            finish = release + 10 + offset
            rows.append(dict(session_id="s%d" % index, request_id="s%d-r%d" % (index, turn), round=turn,
                             session_start_ms=start, release_ms=release, deadline_ms=release + 50,
                             think_ms=0 if turn == 1 else 5, status="completed", finish_ms=finish,
                             llm_start_ms=release + 1, first_token_ms=release + 2,
                             image_ready_ms=release + 4, image_start_ms=release + 6, image_finish_ms=finish,
                             token_intervals_ms=[1, 2], llm_token_ids=[11, 22, 33], token_count=3))
    return dict(status="completed", load=load, seed=301, arm=arm, session_count=2, rows=rows,
                trace_sha256="trace-" + str(load), device_identity={"uuid": "test-device"}, cost_model_sha256="cost",
                elapsed_ms=200 + offset, quality={"passed": True, "semantic_quality": "not_evaluated"})


def matrix():
    return [record(load, arm, index) for load in (.5, .9) for index, arm in enumerate(ARMS)]


class SessionReportingTests(unittest.TestCase):
    def test_sheet_append_does_not_rescan_bounds_and_preserves_roundtrip(self):
        try:
            from openpyxl import Workbook
            from openpyxl.worksheet.worksheet import Worksheet
        except ImportError:
            self.skipTest("optional reports extra not installed")
        book = Workbook()
        book.remove(book.active)
        expected = {}
        # Include ragged rows and formula-like strings; the former triggered
        # whole-sheet column scans when the exporter selected its latest row.
        rows = [[index, "=1+1", "+literal", "@value"] if index % 2 else [index, "-literal"]
                for index in range(256)]
        with patch.object(Worksheet, "max_row", new_callable=PropertyMock,
                          side_effect=AssertionError("row bounds scanned during append")), \
                patch.object(Worksheet, "max_column", new_callable=PropertyMock,
                             side_effect=AssertionError("column bounds scanned during append")):
            _write_sheet(book, "linear", ("index", "first", "second", "third"), iter(rows), expected)
        self.assertEqual([["index", "first", "second", "third"], *rows], expected["linear"])
        self.assertEqual("A1:D257", book["linear"].auto_filter.ref)
        self.assertEqual("A2", book["linear"].freeze_panes)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "linear.xlsx"
            book.save(path)
            book.close()
            _verify_workbook(path, expected)

    def test_sheet_rejects_oversize_literal_before_silent_truncation(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("optional reports extra not installed")
        for rows, error in (([["x" * 32768]], "cell limit"),
                            ([[None] * 16385], "sheet limit")):
            with self.subTest(error=error):
                book = Workbook()
                with self.assertRaisesRegex(ValueError, error):
                    _write_sheet(book, "invalid", ("value",), rows, {})
                book.close()

    def test_paired_inputs_allow_causal_release_differences(self):
        result = evaluate_pilot(matrix(), session_count=2)
        self.assertTrue(result["technical_complete"])
        self.assertEqual(2, len(result["paired_groups"]))
        self.assertIsNone(result["selected_candidate"])

    def test_missing_failed_and_cross_device_pairs_never_qualify(self):
        self.assertFalse(evaluate_pilot(matrix()[:-1], session_count=2)["technical_complete"])
        records = matrix()
        records[0]["quality"]["passed"] = False
        self.assertFalse(evaluate_pilot(records, session_count=2)["technical_complete"])
        records = matrix()
        records[0]["device_identity"] = "another GPU"
        self.assertFalse(evaluate_pilot(records, session_count=2)["technical_complete"])

    def test_fake_second_turn_release_is_rejected(self):
        records = matrix()
        records[0]["rows"][1]["release_ms"] = 1
        self.assertFalse(evaluate_pilot(records, session_count=2)["technical_complete"])

    def test_session_success_requires_both_rounds_on_time(self):
        value = record()
        value["rows"][0]["deadline_ms"] = 5
        result = aggregate([value], session_count=2)
        self.assertEqual(3, result["ontime_rounds"])
        self.assertEqual(1, result["successful_sessions"])
        self.assertEqual(15, result["round_goodput_rps"])
        self.assertEqual(5, result["session_goodput_rps"])
        self.assertEqual(4, result["ttft_ms"]["n"])
        self.assertEqual(8, result["token_interval_ms"]["n"])

    def test_missing_does_not_shrink_planned_denominator(self):
        value = record()
        value["rows"].pop()
        result = aggregate([value], session_count=2)
        self.assertEqual(4, result["planned_rounds"])
        self.assertEqual(1, result["missing_rounds"])
        self.assertEqual(1, result["successful_sessions"])
        self.assertIsNone(aggregate([])["round_goodput_rps"])

    def _root(self, folder):
        root = Path(folder)
        (root / "config.json").write_text(json.dumps(dict(session_count=2, phases=["qualify", "pilot", "state"])))
        for phase in ("qualify", "pilot", "state"):
            (root / phase).mkdir()
            value = dict(status="passed", records=matrix() if phase == "pilot" else [],
                         conclusion="not_applicable" if phase == "state" else "passed")
            (root / phase / "result.json").write_text(json.dumps(value))
        return root

    def test_export_roundtrip_preserves_literals_long_text_and_exact_ids(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("optional reports extra not installed")
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            (root / "unusual.json").write_text(json.dumps({"formula": "=SUM(1,2)", "long": "x" * 40000,
                                                          "id": 98765432101234567890123456789}))
            (root / "pilot" / "execution-events.jsonl").write_text('{"kind":"token","at_ms":1}\n{"kind":"token","at_ms":2}\n')
            (root / "pilot" / "output.npy").write_bytes(b"test bytes; array is never loaded by exporter")
            value = report(root)
            self.assertEqual("verified", value["archive"]["status"])
            files = {item["path"]: item for item in value["archive"]["files"]}
            self.assertTrue(files["pilot/execution-events.jsonl"]["cleanup_eligible"])
            self.assertTrue(files["pilot/output.npy"]["cleanup_eligible"])
            self.assertFalse(files["pilot/result.json"]["cleanup_eligible"])
            book = openpyxl.load_workbook(root / "results.xlsx", data_only=False)
            self.assertEqual(25, book["rounds"].max_row)
            values = list(book["evidence"].values)
            formula = next(row for row in values if row[0] == "unusual.json" and row[1] == "/formula")
            self.assertEqual("=SUM(1,2)", formula[-1])
            self.assertEqual(40000, sum(len(row[-1]) for row in values if row[0] == "unusual.json" and row[1] == "/long"))
            exact_id = next(row for row in values if row[0] == "unusual.json" and row[1] == "/id")
            self.assertEqual("98765432101234567890123456789", exact_id[-1])
            book.close()

    def test_activity_protects_heavy_export_and_retention_survives_refresh(self):
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            ledger = root / "budget.json"
            ledger.write_text(json.dumps({"limit_gpu_seconds": None, "active": {"pid": 1}}))
            value = report(root, ledger)
            self.assertEqual("protected", value["archive"]["status"])
            self.assertFalse((root / "results.xlsx").exists())
            receipt = dict(status="complete", files=[dict(path="gone.npy", sha256="original", bytes=2, state="deleted")])
            (root / "cleanup-summary.json").write_text(json.dumps(receipt))
            self.assertEqual(receipt, report(root, ledger)["retention"])
            self.assertTrue(build_report(root, ledger)["pilot"]["technical_complete"])

    def test_invalid_event_summary_not_cleanable(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("optional reports extra not installed")
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            (root / "bad.jsonl").write_text("not-json\n")
            value = report(root)
            item = next(item for item in value["archive"]["files"] if item["path"] == "bad.jsonl")
            self.assertFalse(item["cleanup_eligible"])

    def test_deleted_event_summary_and_hash_survive_final_refresh(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("optional reports extra not installed")
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            events = root / "actions.jsonl"
            events.write_text('{"action":"llm_decode","start_ms":1,"end_ms":2}\n')
            first = report(root)
            source = next(item for item in first["archive"]["files"] if item["path"] == "actions.jsonl")
            receipt = dict(status="completed", files=[dict(source, state="deleted")])
            events.unlink()
            (root / "cleanup-summary.json").write_text(json.dumps(receipt))
            final = report(root)
            self.assertIn(dict(source, state="deleted", cleanup_eligible=False), final["archive"]["files"])
            book = openpyxl.load_workbook(root / "results.xlsx", read_only=True)
            rows = list(book["events"].values)
            self.assertTrue(any(row[1] == "/kinds/llm_decode" and row[-1] == "1" for row in rows[1:]))
            book.close()
            # Another refresh must retain the same summary, not rely on raw data.
            self.assertEqual("verified", report(root)["archive"]["status"])

    def test_protocol_config_and_ledger_are_used_without_duplicate_config_file(self):
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            (root / "config.json").unlink()
            ledger = root / "ledger.json"
            ledger.write_text('{"active":{"pid":123}}')
            (root / "protocol.json").write_text(json.dumps(dict(config={"session_count": 2}, ledger_path=str(ledger))))
            value = build_report(root)
            self.assertEqual(2, value["config"]["session_count"])
            self.assertTrue(value["pilot"]["technical_complete"])
            self.assertEqual("protected", report(root)["archive"]["status"])

    def test_two_file_archive_recovers_original_evidence_without_reevaluation(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("optional reports extra not installed")
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            # Numeric dictionary keys, escaped pointers and long literals must
            # retain their original containers and values after JSON retirement.
            config = dict(session_count=2, special={"0": {"a/b~c": "=" + "x" * 40000}}, empty=[{}, []])
            (root / "config.json").write_text(json.dumps(config))
            original = report(root)
            for path in root.rglob("*.json"):
                path.unlink()
            self.assertTrue(archived(root))
            restored = load_archive(root)
            self.assertEqual(config, restored["config"])
            self.assertEqual(original["phases"], restored["phases"])
            self.assertEqual(original["pilot"], restored["pilot"])
            before = (root / "results.xlsx").read_bytes()
            self.assertTrue(report(root)["archive_only"])
            self.assertEqual(before, (root / "results.xlsx").read_bytes())
            receipt = dict(status="completed", dependencies_closed=True, final_archive=True,
                           files=[dict(path="config.json", state="deleted", sha256="stored-original")])
            closed = report(root, retention=receipt)
            self.assertEqual(original["pilot"], closed["pilot"])
            self.assertEqual(receipt, load_archive(root)["retention"])
            self.assertEqual(original["phases"], load_archive(root)["phases"])

    def test_partial_control_cleanup_updates_receipt_without_recomputing_missing_phases(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("optional reports extra not installed")
        with tempfile.TemporaryDirectory() as folder:
            root = self._root(folder)
            original = report(root)
            (root / "pilot/result.json").unlink()
            receipt = dict(status="retained", dependencies_closed=True,
                           files=[dict(path="pilot/result.json", state="deleted")])
            result = report(root, retention=receipt, archive_only=True)
            self.assertTrue(result["pilot"]["technical_complete"])
            self.assertEqual(original["phases"], result["phases"])
            self.assertEqual(receipt, load_archive(root)["retention"])

    def test_technically_closed_negative_capacity_results_allow_verified_raw_cleanup(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("optional reports extra not installed")
        for conclusion in ("observed_transient_capacity_failure", "observed_no_viable_state_policy"):
            with self.subTest(conclusion=conclusion), tempfile.TemporaryDirectory() as folder:
                root = self._root(folder)
                state = dict(status="stopped", technical_complete=True, eligible=False, conclusion=conclusion)
                (root / "state/result.json").write_text(json.dumps(state))
                (root / "state/supervisor.json").write_text(json.dumps(dict(state, returncode=0,
                    final_occupancy=dict(status="exclusive_observed", owned=[], foreign=[]))))
                (root / "pilot/output.npy").write_bytes(b"numerically checked image")
                value = report(root)
                item = next(item for item in value["archive"]["files"] if item["path"] == "pilot/output.npy")
                self.assertTrue(item["cleanup_eligible"])
                self.assertEqual(state, value["phases"]["state"])
                self.assertFalse(value["phases"]["state"]["eligible"])

    def test_failed_incomplete_or_unsettled_state_keeps_raw_dependencies(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("optional reports extra not installed")
        for fault in ("failed", "incomplete", "supervisor_failure", "bad_returncode", "cleanup_unverified"):
            with self.subTest(fault=fault), tempfile.TemporaryDirectory() as folder:
                root = self._root(folder)
                state = dict(status="stopped", technical_complete=True, eligible=False)
                supervisor = dict(state, returncode=0)
                if fault == "failed":
                    state["status"] = "failed"
                elif fault == "incomplete":
                    state["technical_complete"] = False
                elif fault == "supervisor_failure":
                    supervisor["supervisor_failure"] = {"reason": "isolation lost"}
                elif fault == "bad_returncode":
                    supervisor["returncode"] = 1
                else:
                    (root / "state/cleanup-unverified.json").write_text("{}")
                (root / "state/result.json").write_text(json.dumps(state))
                (root / "state/supervisor.json").write_text(json.dumps(supervisor))
                (root / "pilot/output.npy").write_bytes(b"retained for investigation")
                value = report(root)
                item = next(item for item in value["archive"]["files"] if item["path"] == "pilot/output.npy")
                self.assertFalse(item["cleanup_eligible"])


if __name__ == "__main__":
    unittest.main()
