import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from ditserve import cli
from ditserve.migration import cli as migration_cli
from ditserve.migration.reporting import write_report


class MigrationCLITest(unittest.TestCase):
    def test_help_is_lazy(self):
        script = "from ditserve.cli import main\nimport sys\ntry: main(['migrate','--help'])\nexcept SystemExit as e: assert e.code == 0\nassert 'torch' not in sys.modules\nassert 'diffusers' not in sys.modules\n"
        result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_archived_entry_has_explicit_message(self):
        stream = io.StringIO()
        with contextlib.redirect_stderr(stream):
            self.assertEqual(cli.main(["benchmark", "--help"]), 2)
        self.assertIn("legacy/README.md", stream.getvalue())

    def test_explicit_download_exits_offline_mode(self):
        def fake(*args, **kwargs):
            self.assertEqual(os.environ["HF_HUB_OFFLINE"], "0")
            return {"status": "verified"}
        with patch.dict(os.environ, {"HF_HUB_OFFLINE": "1"}), \
             patch("ditserve.migration.models.download_models", side_effect=fake), \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(migration_cli.main(["models", "download"]), 0)

    def test_default_check_is_cpu_and_failure_recorded(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch("ditserve.migration.checks.cpu_check", return_value={"status": "failed", "reason": "missing dependency"}), \
             patch("ditserve.migration.checks.gpu_check") as gpu, \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(migration_cli.main(["check", "--output", directory]), 1)
            gpu.assert_not_called()
            self.assertTrue((Path(directory) / "results.xlsx").is_file())
            self.assertIn("failed", (Path(directory) / "REPORT.md").read_text())

    def test_report_preserves_long_strings_ids_and_failures(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            value = {"status": "failed", "missing": None, "large_id": 10**24 + 1,
                     "long": "=literal" * 10000, "unicode": "迁移", "count": 0}
            result = write_report(directory, value)
            book = load_workbook(result["workbook"], read_only=True, data_only=False)
            rows = list(book["evidence"].iter_rows(min_row=2, values_only=True))
            book.close()
            grouped = {}
            for key, part, text in rows:
                grouped.setdefault(key, []).append((part, text))
            rebuilt = {key[1:]: json.loads("".join(v for _, v in sorted(parts)))
                       for key, parts in grouped.items()}
            self.assertEqual(rebuilt, value)
            with self.assertRaises(ValueError):
                write_report(directory, value)

    def test_existing_evidence_prevents_work_before_report(self):
        with tempfile.TemporaryDirectory() as directory:
            (Path(directory) / "REPORT.md").write_text("previous evidence")
            with patch("ditserve.migration.checks.gpu_check") as gpu, \
                 contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(migration_cli.main(["check", "--suite", "gpu", "--output", directory]), 1)
                gpu.assert_not_called()

    def test_stale_model_paths_revisions_and_interpreter_fail_before_gpu(self):
        from ditserve.migration.configure import environment_config
        uuid = "GPU-00000000-0000-0000-0000-000000000001"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "environment.json"
            for key, wrong in (("path", "/home/old-machine/models/wan"),
                               ("revision", "old-revision"), ("python", "/old/venv/bin/python")):
                value = environment_config("/models", "/data", uuid)
                if key == "python":
                    value[key] = wrong
                else:
                    value["models"]["wan"][key] = wrong
                path.write_text(json.dumps(value))
                with patch("ditserve.migration.models.verify_models") as verify, \
                     patch("ditserve.research.preflight.discover_gpus") as discover:
                    with self.assertRaises(ValueError):
                        migration_cli.smoke(path, ("core",))
                    verify.assert_not_called()
                    discover.assert_not_called()

    def test_longlive_invalid_cpu_evidence_blocks_gpu_work(self):
        from ditserve.migration.configure import environment_config
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "environment.json"
            path.write_text(json.dumps(environment_config("/models", "/data",
                "GPU-00000000-0000-0000-0000-000000000001")))
            with patch("ditserve.migration.models.verify_models", return_value={"status": "verified"}), \
                 patch("ditserve.directions.runner.model_inventory", return_value={"available": False}), \
                 patch("ditserve.migration.checks.run_owned") as run, \
                 patch("ditserve.research.preflight.discover_gpus") as discover:
                result = migration_cli.smoke(path, ("longlive",))
            self.assertEqual(result["status"], "failed")
            self.assertFalse(result["gpu_work_started"])
            run.assert_not_called()
            discover.assert_not_called()


if __name__ == "__main__":
    unittest.main()
