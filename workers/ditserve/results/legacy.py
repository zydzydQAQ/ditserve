"""Read-only compatibility for preserved experiment archives.

Verify historical gzip evidence, workbook cell types/values and report text
without rebuilding, publishing or cleaning any archived files. Original hash
and rendering rules remain intact; new archives use the unified results API.
"""

import gzip
import hashlib
import json
import math
from pathlib import Path
import re


NOTICE = ("本归档保留表格、报告和小型历史记录，不包含原始数组、latent 或完整 trace。"
          "合成结果不是真机性能，参考一致性不等于感知质量，未知质量不计入 SLO 成功。")
TABLE_NAMES = ("Config", "Requests", "Results", "Files")
PURPOSE = "比较固定请求流下各调度策略的参考行为。"
METHOD = "使用保存的 workload、成本模型和全部到达请求，比较完成、超时、SLO 与参考状态一致性。"


def _flatten(value, prefix="", lists=False):
    result = {}
    for key, child in value.items():
        name = prefix + key
        if isinstance(child, dict):
            result.update(_flatten(child, name + ".", lists))
        elif lists or not isinstance(child, list):
            result[name] = child
    return result


def _unit(name):
    if name == "goodput_per_sec":
        return "requests/s"
    if name.endswith(("_ms", ".wall_ms", "_wall_ms_sum")):
        return "ms"
    if name.endswith(("_seconds", "_sec")):
        return "s"
    if name.endswith(("_bytes", ".bytes")):
        return "bytes"
    if name.endswith("_mib"):
        return "MiB"
    if "_over_" in name:
        return "ratio"
    return ""


def _cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) > 2 ** 53:
        return str(value)
    return value


def _columns(rows):
    return list(dict.fromkeys(key for row in rows for key in row))


def _tables(data):
    tables = {name: [] for name in TABLE_NAMES}
    summaries = [(name, file) for name, file in data["files"].items()
                 if Path(name).name == "summary.json" and isinstance(json.loads(file["text"]).get("runs"), dict)]
    selected_summary = min(summaries, key=lambda item: len(Path(item[0]).parts))[0] if summaries else None
    for name, file in data["files"].items():
        tables["Files"].append({"source": name, "bytes": file["bytes"], "sha256": file["sha256"]})
        if not name.endswith(".json"):
            continue
        value = json.loads(file["text"])
        if Path(name).name == "workload.json" and isinstance(value, list):
            tables["Requests"].extend(dict(source=name, **row) for row in value)
        if not isinstance(value, dict):
            continue
        config = value.get("config")
        if isinstance(config, dict):
            tables["Config"].extend({"source": name, "parameter": key, "value": item, "unit": _unit(key)}
                                    for key, item in _flatten(config, lists=True).items())
        for key in ("cost_model", "core"):
            if isinstance(value.get(key), dict):
                parameters = _flatten({key: value[key]}, lists=True)
                tables["Config"].extend({"source": name, "parameter": field, "value": item, "unit": _unit(field)}
                                        for field, item in parameters.items())
        if Path(name).name == "cost_model.json":
            tables["Config"].extend({"source": name, "parameter": "cost_model." + key, "value": item, "unit": _unit(key)}
                                    for key, item in _flatten(value, lists=True).items())
        if name == selected_summary:
            tables["Results"].extend(dict(source=name, case=case, **row) for case, row in value["runs"].items())
    return tables


def _markdown_table(rows, fields):
    if not rows or not fields:
        return []
    lines = ["| " + " | ".join(fields) + " |", "| " + " | ".join("---" for _ in fields) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(row.get(field, "") if row.get(field) is not None else "").replace("|", "\\|").replace("\n", " ") for field in fields) + " |")
    return lines


def _conclusion(data, tables):
    rows = tables["Results"]
    observed = [row["parity_failures"] for row in rows if row.get("parity_failures") is not None]
    failures = sum(observed) if observed else "未知"
    return ("保存 %d 条策略/配置结果，参考状态一致性失败计数为 %s。"
            "时间、资源和 goodput 均来自合成成本；它们不能作为真实 GPU 性能或感知质量证据。" % (len(rows), failures))


def _report(data, tables):
    rows = tables["Results"]
    statuses = {}
    for row in rows:
        status = row.get("status") or "synthetic"
        statuses[status] = statuses.get(status, 0) + 1
    lines = ["# " + data["title"], "", "## 目的与方法", "", PURPOSE, "", METHOD, "",
             "工作簿：[results.xlsx](results.xlsx)。可重建数据：[data.json.gz](data.json.gz)。", "",
             "Config 表保存逐实验配置；Results 保存逐策略汇总。"
             "字段后缀 `_ms`、`_seconds`、`_bytes`、`_mib` 分别表示毫秒、秒、字节、MiB；`goodput_per_sec` 为请求/秒。", "",
             "## 实际结果", "", _conclusion(data, tables), "",
             "保存 %d 份原始小文件、%d 条结果记录；另有 %d 份二进制产物仅登记哈希、不嵌入归档。" %
             (len(data["files"]), len(rows), len(data.get("external_artifacts", []))),
             "状态计数：`%s`。空白测项表示未知或不适用，不填零。" % json.dumps(statuses, ensure_ascii=False), ""]
    fields = [field for field in ("case", "completed", "rejected", "timed_out", "slo_attainment", "goodput_per_sec",
                                  "p99_latency_ms", "mean_batch_size", "preemptions", "offloads", "restores", "parity_failures")
              if any(field in row for row in rows)]
    lines.extend(_markdown_table(rows, fields))
    lines.extend(["", "## 结论边界与保留策略", "", NOTICE, "",
                  "原始 JSON 与报告文本原样存于压缩数据中；重建报告无需 GPU、模型或原始数组。", ""])
    return "\n".join(lines)


def _check_links(path):
    for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
        if "://" in target or target.startswith("#"):
            continue
        if not (path.parent / target).is_file():
            raise ValueError("broken report link: " + target)


def _check_cells(sheet, expected):
    if sheet.max_row != len(expected) or sheet.max_column != max(map(len, expected)):
        raise ValueError("workbook dimensions differ: " + sheet.title)
    for row, values in zip(sheet, expected):
        for index, cell in enumerate(row):
            wanted = _cell(values[index]) if index < len(values) else None
            if isinstance(wanted, str):
                wanted = wanted[:32767] or None
            actual = cell.value
            if wanted is None:
                valid = actual is None
            elif isinstance(wanted, bool):
                valid = isinstance(actual, bool) and actual == wanted and cell.data_type == "b"
            elif isinstance(wanted, (int, float)):
                valid = (isinstance(actual, (int, float)) and not isinstance(actual, bool) and
                         cell.data_type == "n" and math.isclose(actual, wanted, rel_tol=1e-14, abs_tol=1e-12))
            else:
                valid = isinstance(actual, str) and actual == wanted and cell.data_type == "s"
            if not valid:
                raise ValueError("workbook cell differs: " + sheet.title + "!" + cell.coordinate)


def _verify_step(root, key):
    from openpyxl import load_workbook

    step = Path(root) / key
    with gzip.open(step / "data.json.gz", "rt", encoding="utf-8") as stream:
        data = json.load(stream)
    for name, record in data["files"].items():
        encoded = record["text"].encode("utf-8")
        if len(encoded) != record["bytes"] or hashlib.sha256(encoded).hexdigest() != record["sha256"]:
            raise ValueError("archived evidence changed: " + name)
    tables = _tables(data)
    book = load_workbook(step / "results.xlsx", read_only=True, data_only=False)
    try:
        overview = [["title", data["title"]], ["scope", NOTICE], ["raw_artifacts_available", False],
                    ["historical_source_files", len(data["files"])]]
        overview.extend([name, _cell(value)] for name, value in data.get("counts", {}).items())
        _check_cells(book["Overview"], overview)
        for name, rows in tables.items():
            columns = _columns(rows) or ["no_records"]
            _check_cells(book[name], [columns] + [[row.get(column) for column in columns] for row in rows])
    finally:
        book.close()
    if (step / "REPORT.md").read_text(encoding="utf-8") != _report(data, tables):
        raise ValueError("report differs from archived data: " + key)
    _check_links(step / "REPORT.md")
    return data


def verify(root):
    """Verify evidence bytes, every workbook cell and report links."""
    root = Path(root).resolve()
    manifest = json.loads((root / "archive.json").read_text(encoding="utf-8"))
    data = {key: _verify_step(root, key) for key in manifest["steps"]}
    _check_links(root / "INDEX.md")
    return {"status": "verified", "steps": len(data),
            "source_files": sum(len(step["files"]) for step in data.values()),
            "raw_artifacts_available": False}


def _load_session(book):
    from . import _json, _restore

    def restore(rows):
        converted = []
        for pointer, kind, part, scalar in rows:
            pointer = pointer or ""
            if kind in ("dict", "list", "empty_dict", "empty_list"):
                converted.append((pointer, kind[6:] if kind.startswith("empty_") else kind, 0, "0"))
            else:
                # Old scalar fragments join before quoting string values.
                converted.append((pointer, kind, int(part), scalar))
        groups = {}
        for pointer, kind, part, scalar in converted:
            groups.setdefault((pointer, kind), []).append((part, scalar))
        child_counts = {}
        for pointer, kind in groups:
            if pointer:
                parent = pointer.rsplit("/", 1)[0]
                child_counts[parent] = child_counts.get(parent, 0) + 1
        nodes = []
        for (pointer, kind), parts in groups.items():
            parts.sort()
            if [p for p, v in parts] != list(range(len(parts))):
                raise ValueError("invalid legacy fragment sequence")
            if kind in ("dict", "list"):
                raw = str(child_counts.get(pointer, 0))
            elif kind == "str":
                raw = _json("".join(str(v) if v is not None else "" for p, v in parts))
            elif kind == "bool":
                if len(parts) != 1 or type(parts[0][1]) is not bool:
                    raise ValueError("invalid legacy boolean")
                raw = _json(parts[0][1])
            elif kind == "null":
                raw = "null"
            else:
                raw = "".join(str(v) for p, v in parts)
            nodes.append((pointer, kind, 0, raw))
        return _restore(nodes)
    summary = restore(book["summary"].iter_rows(min_row=2, values_only=True))
    if summary.get("schema_version") != "ditserve.sessions.report.v1":
        raise ValueError("unrecognized legacy two-file workbook")
    result = {"summary": summary}
    for name in ("config", "budget", "retention", "provenance"):
        result[name] = restore(book[name].iter_rows(min_row=2, values_only=True))
    evidence = {}
    for row in book["evidence"].iter_rows(min_row=2, values_only=True):
        evidence.setdefault(row[0], []).append(row[1:])
    result["evidence"] = {name: restore(rows) for name, rows in evidence.items()}
    result["tables"] = {sheet.title: list(sheet.values) for sheet in book if sheet.title not in {
        "summary", "config", "budget", "retention", "provenance", "evidence"}}
    return result


def _load_wan_profile(book):
    """Read the immutable, typed path-array format of the original Wan pilot."""
    from . import _restore

    headers = ("path_json", "type", "part", "value_json_fragment")
    def decode(name):
        if tuple(next(book[name].values)) != headers:
            raise ValueError("invalid legacy Wan worksheet header")
        nodes, order = [], {}
        for path, kind, part, scalar in book[name].iter_rows(min_row=2, values_only=True):
            segments = json.loads(path)
            if not isinstance(segments, list) or any(type(x) not in (str, int) for x in segments):
                raise ValueError("invalid legacy Wan path")
            pointer = "".join("/" + str(x).replace("~", "~0").replace("/", "~1") for x in segments)
            nodes.append((pointer, kind, part, scalar))
            if segments:
                keys = order.setdefault(tuple(segments[:-1]), [])
                if segments[-1] not in keys:
                    keys.append(segments[-1])
        def ordered(value, path=()):
            if isinstance(value, dict):
                return {key: ordered(value[key], path + (key,)) for key in order.get(path, [])}
            if isinstance(value, list):
                return [ordered(item, path + (index,)) for index, item in enumerate(value)]
            return value
        return ordered(_restore(nodes))
    summary, value = decode("summary"), decode("evidence")
    if value.get("schema_version") != "ditserve.wan_breakdown.v1":
        raise ValueError("unsupported legacy Wan schema")
    # This historical writer hashed insertion order, unlike the new format.
    encoded = json.dumps(value, ensure_ascii=True, separators=(",", ":"), allow_nan=False)
    if hashlib.sha256(encoded.encode("ascii")).hexdigest() != summary.get("payload_sha256"):
        raise ValueError("legacy Wan payload hash mismatch")
    return value


def load_book(book):
    """Read original typed workbooks without rebuilding or changing evidence."""
    for sheet in book:
        if any(cell.data_type == "f" for row in sheet for cell in row):
            raise ValueError("legacy workbook contains a formula")
    if "evidence" in book and tuple(next(book["evidence"].values)) == ("path_json", "type", "part", "value_json_fragment"):
        return _load_wan_profile(book)
    return _load_session(book)
